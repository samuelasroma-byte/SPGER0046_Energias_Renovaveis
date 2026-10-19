"""Revise B24/R01 from official metadata while preserving the original Word."""
from pathlib import Path
import csv
import json
import re

from docx import Document
from docx.shared import Cm, Pt

ROOT=Path(__file__).resolve().parents[1]
DATA=json.loads((ROOT/'references/pvdaq_dataset_citacao.json').read_text())
DEST=ROOT/'bibliografia/atualizada'
SOURCE=ROOT/'bibliografia/outputs/Referencias_por_Subsecao.docx'


def main():
    DEST.mkdir(exist_ok=True)
    doc=Document(SOURCE)
    paragraph=next(p for p in doc.paragraphs if '10.25984/1846021' in p.text)
    old=paragraph.text
    # Preserve paragraph formatting and the title emphasis of the original.
    for child in list(paragraph._p):
        if child.tag.endswith('}pPr'):continue
        paragraph._p.remove(child)
    citation=DATA['citation_abnt'];before,title,after=citation.partition(DATA['title'])
    for text,bold in [(before,False),(title,True),(after,False)]:
        run=paragraph.add_run(text);run.bold=bold;run.font.name='Times New Roman';run.font.size=Pt(12)
    paragraphs=doc.paragraphs
    index=next(i for i,p in enumerate(paragraphs) if p.text.startswith('REFERÊNCIA B24'))
    heading=paragraphs[index]
    for run in heading.runs:
        run.text=run.text.replace('Documento técnico','Conjunto de dados')
    block=[p._p for p in paragraphs[index:index+3]]
    anchor_index=next(i for i,p in enumerate(paragraphs) if p.text.startswith('REFERÊNCIA B08'))
    anchor=paragraphs[anchor_index+2]._p
    # Keep the alphabetical list with the new authorship and stable B24 links.
    for element in block:anchor.addnext(element);anchor=element
    anchor=next(p for p in doc.paragraphs if p.text.startswith('Quadro 1'))
    p=anchor.insert_paragraph_before(
        'Revisão de 19 de outubro de 2026. B24 identifica o conjunto PVDAQ/OEDI de Deline e colaboradores (2021); '
        'B27 documenta o acesso. O estudo utiliza o sistema 2107 entre abril e junho de 2024. '
        'O ano da referência é o de publicação do conjunto, enquanto 2024 identifica as observações analisadas. '
        'Os cinco arquivos originais, suas URLs e seus hashes estão em data/raw/fontes_pvdaq.json.')
    p.paragraph_format.first_line_indent=Cm(0)
    p.paragraph_format.space_after=Pt(12)
    for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(12)
    doc.core_properties.subject='39 referências por subseção com origem PVDAQ identificada e revisão B24 de 19/10/2026'
    doc.save(DEST/'Referencias_por_Subsecao_Atualizadas.docx')

    path=ROOT/'bibliografia/principal/Referencias_B01_B39.json'
    rows=json.loads(path.read_text());row=next(r for r in rows if r['code']=='B24')
    row.update(type='Conjunto de dados',reference=citation,
               observacao='Fonte primária OEDI; sistema 2107 e arquivos de 2024 identificados em data/raw/fontes_pvdaq.json',
               source_metadata='references/pvdaq_dataset_citacao.json')
    path.write_text(json.dumps(rows,ensure_ascii=False,indent=2)+'\n')
    path=ROOT/'bibliografia/principal/Correspondencia_B_A.csv'
    with path.open(encoding='utf-8-sig',newline='') as handle:
        reader=csv.DictReader(handle);fields=reader.fieldnames;correspondence=list(reader)
    for r in correspondence:
        if r['codigo_B']=='B24':r.update(tipo=row['type'],referencia_ABNT=citation,observacao=row['observacao'])
    with path.open('w',encoding='utf-8-sig',newline='') as handle:
        writer=csv.DictWriter(handle,fields);writer.writeheader();writer.writerows(correspondence)

    for path in [ROOT/'bibliografia/principal/README.md',*ROOT.glob('secoes/*/referencias/README.md')]:
        text=path.read_text()
        text=text.replace('**Documento técnico.** '+old,'**Conjunto de dados.** '+citation)
        text=text.replace('— Documento técnico | '+old,'— Conjunto de dados | '+citation)
        path.write_text(text)

    bibliography=(ROOT/'bibliografia/outputs/referencias_fundamentacao.bib').read_text()
    pattern=r'@\w+\{([^,]+),\n(?:(?!\n\}).)*10\.25984/1846021(?:(?!\n\}).)*\n\}'
    match=re.search(pattern,bibliography,re.S);assert match
    fields={'title':DATA['title'],'author':' and '.join(DATA['authors']),'year':'2021',
            'publisher':'NREL','howpublished':DATA['repository'],'doi':DATA['doi'],
            'url':DATA['url'],'urldate':DATA['accessed']}
    replacement='@misc{'+match[1]+',\n'+',\n'.join('  '+key+' = {'+value+'}' for key,value in fields.items())+'\n}'
    bibliography=bibliography[:match.start()]+replacement+bibliography[match.end():]
    (DEST/'referencias_fundamentacao.bib').write_text(bibliography)
    ris=(ROOT/'bibliografia/outputs/referencias_fundamentacao.ris').read_text()
    blocks=ris.split('ER  - ');found=0
    for i,block in enumerate(blocks):
        if DATA['doi'] in block:
            found+=1
            blocks[i]='\n\nTY  - DATA\nTI  - '+DATA['title']+'\n'+''.join('AU  - '+author+'\n' for author in DATA['authors'])+'PY  - 2021\nDA  - 2021/12/21\nPB  - NREL\nDB  - '+DATA['repository']+'\nDO  - '+DATA['doi']+'\nUR  - '+DATA['url']+'\nY2  - 2026/10/19\n'
    assert found==1
    (DEST/'referencias_fundamentacao.ris').write_text('ER  - '.join(blocks))
    assert len(rows)==39
    print('B24 e Word revisados; B01–B39 e o Word original foram preservados.')


if __name__=='__main__':main()
