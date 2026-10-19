"""Complete, provenance-aware inventory for writing the four chapter sections."""
from pathlib import Path
import hashlib,io,json,re,sys,zipfile
import numpy as np
import pandas as pd
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
import gerar_documentos as word

UNIT_MODULE={'STC':'W','PTC':'W','A_c':'m²','Length':'m','Width':'m','N_s':'células',
 'I_sc_ref':'A','V_oc_ref':'V','I_mp_ref':'A','V_mp_ref':'V','alpha_sc':'A/K','beta_oc':'V/K',
 'T_NOCT':'°C','a_ref':'V','I_L_ref':'A','I_o_ref':'A','R_s':'Ω','R_sh_ref':'Ω','Adjust':'%','gamma_r':'%/K'}
UNIT_INV={'Vac':'V','Pso':'W','Paco':'W','Pdco':'W','Vdco':'V','C0':'1/W','C1':'1/V','C2':'1/V','C3':'1/V',
          'Pnt':'W','Vdcmax':'V','Idcmax':'A','Mppt_low':'V','Mppt_high':'V'}
UNIT_ASSUMPTIONS={'faiman_u0':'W/(m² K)','faiman_u1':'W s/(m³ K)','albedo':'1',
 'soiling_fraction':'fração','dc_wiring_fraction':'fração','dc_mismatch_fraction':'fração','dc_connections_fraction':'fração',
 'ac_wiring_fraction':'fração','inverter_eta_nom':'fração','inverter_eta_ref':'fração','inverter_capacity_baseline_w_each':'W',
 'nasa_hour_label_offset_hours':'h','illustrative_modules_in_series':'módulos','illustrative_parallel_strings':'strings'}
UNIT_BIPV={'area_m2':'m²','reference_efficiency':'fração','gamma_pdc_per_C':'1/°C',
 'rho_solar':'fração','tau_solar':'fração','tau_visible':'fração','inward_heat_fraction':'fração',
 'faiman_u0':'W/(m² K)','faiman_u1':'W s/(m³ K)','tilt_deg':'°','azimuth_deg':'°'}
LABEL={'faiman_u0':'Coeficiente térmico U₀','faiman_u1':'Coeficiente térmico U₁','albedo':'Albedo',
 'soiling_fraction':'Sujidade adicional de referência','dc_wiring_fraction':'Perda na fiação DC',
 'dc_mismatch_fraction':'Perda por mismatch DC','dc_connections_fraction':'Perda nas conexões DC',
 'ac_wiring_fraction':'Perda na fiação AC','inverter_eta_nom':'Eficiência nominal assumida do inversor',
 'inverter_eta_ref':'Eficiência de referência da equação PVWatts',
 'inverter_capacity_baseline_w_each':'Limite AC assumido por inversor',
 'illustrative_modules_in_series':'Módulos em série no equivalente didático',
 'illustrative_parallel_strings':'Strings paralelas no equivalente didático',
 'area_m2':'Área do exemplo semitransparente','reference_efficiency':'Eficiência elétrica de referência',
 'gamma_pdc_per_C':'Coeficiente de potência','rho_solar':'Refletância solar',
 'tau_solar':'Transmitância solar','tau_visible':'Transmitância visível',
 'inward_heat_fraction':'Fração do calor absorvido dirigida ao interior','tilt_deg':'Inclinação','azimuth_deg':'Azimute'}

def flatten(value,prefix=''):
    rows=[]
    if isinstance(value,dict):
        for k,v in value.items():rows.extend(flatten(v,prefix+'.'+k if prefix else k))
    elif isinstance(value,list):
        for i,v in enumerate(value):rows.extend(flatten(v,f'{prefix}[{i}]'))
    else:rows.append({'metadata_path':prefix,'value':value,'value_type':type(value).__name__})
    return rows

def inventory():
    cfg=json.loads((ROOT/'config/projeto.json').read_text())
    sources=json.loads((ROOT/'data/raw/fontes_pvdaq.json').read_text())
    z=zipfile.ZipFile(ROOT/'data/raw/pvdaq_2107_2024_originais.zip')
    for source in sources:
        assert hashlib.sha256(z.read(source['file'])).hexdigest()==source['sha256'],source['file']
    meta=json.loads(z.read('2107_system_metadata.json'))
    out=ROOT/'outputs/dossie_tecnico';out.mkdir(exist_ok=True)
    (out/'metadados_originais_integrais.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2))
    pd.DataFrame(flatten(meta)).to_csv(out/'metadados_todos_os_campos.csv',index=False)
    files=[];channels=[]
    for name in z.namelist():
        if not name.endswith('.csv'):continue
        frame=pd.read_csv(z.open(name));dates=pd.to_datetime(frame['measured_on'],errors='coerce')
        in_q=(dates>=pd.Timestamp(cfg['period']['start_local']))&(dates<pd.Timestamp(cfg['period']['end_exclusive_local']))
        stamps=dates.dropna().drop_duplicates().sort_values()
        diff=stamps.diff().dropna().dt.total_seconds()/60
        cadence=float(diff.mode().iloc[0])
        files.append(dict(file=name,rows=len(frame),channels=len(frame.columns)-1,
           raw_local_first=str(dates.min()),raw_local_last=str(dates.max()),
           most_frequent_step_minutes=cadence,invalid_timestamps=int(dates.isna().sum()),
           repeated_local_timestamps=int(dates.duplicated().sum()),quarter_rows=int(in_q.sum()),
           convention='horário local declarado PST8PDT; interpretação America/Los_Angeles'))
        for column in frame.columns:
            if column=='measured_on':continue
            metadata_key=column
            metric=meta['Metrics'].get(column,{})
            matching='nome exato'
            if not metric:
                metric_id=int(column.rsplit('_',1)[1])
                aliases=[(key,item) for key,item in meta['Metrics'].items() if item.get('metric_id')==metric_id]
                assert len(aliases)==1,('Canal sem metadado inequívoco',column)
                metadata_key,metric=aliases[0]
                matching='ID único; divergência no nome do cadastro'
            series=pd.to_numeric(frame[column],errors='coerce');q=series.loc[in_q]
            channel=dict(file=name,column=column,metadata_key=metadata_key,metadata_matching=matching,metric_id=metric.get('metric_id'),
              common_name=metric.get('common_name'),sensor_name=metric.get('sensor_name'),
              raw_unit=metric.get('raw_units'),stored_unit=metric.get('units'),
              aggregation=metric.get('aggregation_type'),source_type=metric.get('source_type'),
              scale=metric.get('calc_scale'),offset=metric.get('calc_offset'),offline=metric.get('offline'),
              annual_rows=len(frame),annual_valid=int(series.notna().sum()),annual_missing=int(series.isna().sum()),
              quarter_rows=int(in_q.sum()),quarter_valid=int(q.notna().sum()),quarter_missing=int(q.isna().sum()),
              quarter_minimum=q.min(),quarter_mean=q.mean(),quarter_maximum=q.max(),
              quarter_negative=int(q.lt(0).sum()),quarter_zero=int(q.eq(0).sum()))
            if column.startswith('poa'):channel['role']='entrada radiativa medida'
            elif column.startswith(('ambient','wind_speed')):channel['role']='entrada meteorológica medida'
            elif column.startswith('wind_direction'):channel['role']='disponível; não usado na cadeia principal'
            elif '_ac_power_' in column or column.startswith('meter_'):channel['role']='referência elétrica observada; não entrada do modelo'
            else:channel['role']='diagnóstico elétrico adicional; não entrada radiativa ou térmica'
            channels.append(channel)
    assert len(channels)==125==len(meta['Metrics'])
    assert {c['metric_id'] for c in channels}=={m['metric_id'] for m in meta['Metrics'].values()}
    pd.DataFrame(files).to_csv(out/'inventario_arquivos.csv',index=False)
    pd.DataFrame(channels).to_csv(out/'dicionario_completo_125_canais.csv',index=False)
    for key,name in [('Inverters','inversores'),('Modules','grupos_de_modulos'),('Mount','montagem'),
                     ('Meters','medidores'),('Other Instruments','instrumentos')]:
        pd.DataFrame([{'metadata_label':k,**v} for k,v in meta[key].items()]).to_csv(out/(name+'.csv'),index=False)
    all_config=[{'section':section,**row} for section in ['site','period','system','assumptions','bipv','workflow']
                for row in flatten(cfg[section],section)]
    pd.DataFrame(all_config).to_csv(out/'configuracao_completa_modelo.csv',index=False)
    return cfg,meta,pd.DataFrame(files),pd.DataFrame(channels),out

def main():
    cfg,meta,files,channels,out=inventory()
    module=json.loads((ROOT/'config/modulo_cec.json').read_text());inv=json.loads((ROOT/'config/inversor_cec.json').read_text())
    summary=json.loads((ROOT/'outputs/resumo_execucao.json').read_text())
    csv=ROOT/'outputs/tabelas/csv';descriptive=pd.read_csv(csv/'estatistica_descritiva.csv')
    doc=word.new_document('Dossiê técnico dos dados e da instalação fotovoltaica',
       'Base para a redação compartilhada do capítulo Recurso solar e modelagem da conversão fotovoltaica, '
       'SPGER0046 da UFPB. O documento reúne a instalação Farm Solar Array, os arquivos do PVDAQ 2107, '
       'as variáveis de entrada, as referências elétricas e os parâmetros do cálculo, organizados pelas quatro seções. '
       'O ano 2026 na fonte indica a elaboração desta compilação; as séries analisadas são de 2024.')
    word.paragraph(doc,'Há 125 canais medidos nos quatro CSV originais, além do JSON de metadados. '
       'Os apêndices digitais preservam todos os campos de metadados e as estatísticas de cada canal. '
       'As tabelas deste dossiê descrevem os dados; as séries completas permanecem nos arquivos originais.')
    nums={'Tabela':0,'Quadro':0}
    def tab(title,headers,rows,widths=None,grid=False,note=''):
        kind='Quadro' if grid else 'Tabela';nums[kind]+=1
        before=len(doc.tables)
        t=word.table(doc,f'{kind} {nums[kind]} — {title}',headers,rows,widths,grid,note)
        for part in doc.tables[before:]:
            for index,row in enumerate(part.rows):
                for cell in row.cells:
                    for margin in cell._tc.xpath('./w:tcPr/w:tcMar/w:top|./w:tcPr/w:tcMar/w:bottom'):
                        margin.set(qn('w:w'),'35')
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
                        p.paragraph_format.line_spacing=1.0
                        following=part._tbl.getnext()
                        has_source=following is not None and ''.join(following.itertext()).startswith('Fonte:')
                        p.paragraph_format.keep_with_next=index<len(part.rows)-1 or has_source
        return t
    def head(title,new=False):
        p=doc.add_paragraph(title,'Heading 1');p.paragraph_format.keep_with_next=True
        p.paragraph_format.page_break_before=new
        for r in p.runs:r.bold=True;r.font.size=Pt(12)
    def page():
        """Table blocks flow continuously; section headings alone start pages."""
        return None
    def v(x):
        if x is None:return 'Não informado'
        if isinstance(x,bool):return 'Sim' if x else 'Não'
        return f'{x:.12g}'.replace('.',',') if isinstance(x,float) else str(x)
    head('1 Recurso solar localização e séries observacionais')
    tab('Identificação e localização da instalação',['Campo','Informação'],[
       ['Nome público','Farm Solar Array'],['Sistema e sítio PVDAQ','2107; site_id 7211, FSA_1'],
       ['Localidade','Arbuckle, Califórnia, Estados Unidos'],['Clima no cadastro','Csa'],
       ['Fuso no cadastro','PST8PDT'],['Fuso interpretado no cálculo','America/Los_Angeles'],
       ['Origem','PVDAQ / OEDI, dados públicos do Departamento de Energia dos EUA'],
       ['Conteúdo dos arquivos','2024; as datas efetivas constam do inventário'],
       ['Data de início no cadastro do sistema','01/01/2017'],['Data inicial nos grupos de módulos','01/11/2017']],
       [5.2,12.5],True,'As datas cadastradas de início são diferentes; o período de estudo depende dos CSV de 2024.')
    tab('Coordenadas e janela de análise',['Grandeza','Valor','Unidade'],[
       ['Latitude',v(cfg['site']['latitude']),'° N'],['Longitude',v(cfg['site']['longitude']),'° E, negativa a oeste'],
       ['Altitude',10,'m'],['Dias do trimestre',91,'dias'],['Passo comum',15,'min'],
       ['Intervalos totais esperados',8736,'intervalos'],['Elevação solar mínima do filtro',10,'°']],
       [8.1,4.3,5.3],note='De 01/04/2024 00:00 local até 01/07/2024 00:00 local, com fim exclusivo.')
    page()
    tab('Arquivos observacionais e extensão efetivamente disponível',['Arquivo','Registros','Canais','Passo modal (min)'],
       [[r.file,str(r.rows),str(r.channels),v(r.most_frequent_step_minutes)] for _,r in files.iterrows()],
       [10.4,2.4,1.8,3.1],note='Números do arquivo inteiro. O passo modal não garante regularidade ou ausência de lacunas. Limites temporais completos, duplicatas e amostras do trimestre constam do CSV inventario_arquivos.')
    tab('Limites locais das séries originais',['Grupo','Primeiro carimbo','Último carimbo'],
       [[r.file.replace('2107_','').replace('_data_2024.csv',''),r.raw_local_first,r.raw_local_last] for _,r in files.iterrows()],
       [5.7,6,6],True,'O JSON do cadastro termina em 2023; esse campo histórico não descreve a extensão dos CSV de 2024.')
    page()
    tab('Irradiância e geometria disponíveis para a seção 1',['Elemento','Registro e utilização'],[
       ['POA medida','Canal poa_irradiance_o_149574; metric_id 149574; W/m²; agregação original avg.'],
       ['Geometria do sensor','POA se refere ao plano do arranjo. Não há cadastro independente de marca, modelo, calibração ou incerteza do sensor.'],
       ['GHI, DNI e DHI medidas','Não disponíveis neste conjunto. As componentes reconstruídas são estimativas exploratórias.'],
       ['Ângulos solares','SPA para zênite, elevação e azimute; calculados pelas coordenadas e instantes, não sensores angulares.'],
       ['Irradiância extraterrestre','Método Spencer implementado no pvlib; calculada.'],
       ['Carimbo temporal','Tratado como início do intervalo; instante geométrico em +7,5 min nas médias de 15 min e +30 min nas médias horárias. A convenção de início é uma hipótese do tratamento, sem confirmação instrumental independente.'],
       ['Agregação','POA de 5 min em médias de 15 min, exigindo três amostras; não preencher lacunas.']],
       [4.5,13.2],True,'Subseções 1.1 a 1.4; referências de método R03–R06 e R14–R18. A inversão não resolve de modo único e independente três componentes a partir de uma medição.')
    head('2 Meteorologia ambiente e montagem',True)
    tab('Canais meteorológicos e conversões',['Variável','ID','Unidade original','Unidade no modelo'],[
       ['Temperatura do ar',149575,'°F','°C'],['Velocidade do vento',149576,'mph','m/s'],['Direção do vento',149577,'grau','grau']],
       [6.8,2,4.3,4.6],note='T°C=(T°F−32)×5/9; vento em m/s=vento em mph×0,44704. Direção está disponível, mas não entra no modelo principal.')
    tab('Instrumentação meteorológica e representatividade',['Campo','Informação'],[
       ['Estação no cadastro','ws_fss_1; instrument_id 5294; tipo weather station.'],
       ['Fabricante e modelo','Fabricante Unknown; modelo vazio.'],['Altura dos sensores e anemômetro','Não informada. Não se assumiu medição a 10 m.'],
       ['Calibração e incerteza','Não fornecidas no cadastro.'],['Tipo de agregação declarado','Temperatura, velocidade e direção: sample. A coincidência temporal não torna esses registros médias de todo o intervalo.'],
       ['Localização espacial','Meteorologia associada ao sistema PVDAQ; não foi utilizada a estação INMET a 12 km nem uma estação da Paraíba.'],
       ['Dados adicionais ausentes','Temperatura medida da célula, albedo medido, sujidade, cronograma de limpeza, perfil de horizonte e geometria detalhada das fileiras.']],
       [5.1,12.6],True)
    page()
    env=descriptive.iloc[:4]
    tab('Descrição das condições na mesma janela diurna pareada',['Variável','Média','Mínimo','Máximo','Unidade'],
       [[r.variable,word.value(r['mean']),word.value(r['minimum']),word.value(r['maximum']),r.unit] for _,r in env.iterrows()],
       [7.7,2.3,2.3,2.3,3.1],note='4.401 pares de 15 min; a temperatura operacional é calculada, não medida. As estatísticas completas e quartis estão no documento de tabelas.')
    tab('Montagem e entorno',['Campo','Informação'],[
       ['Montagem','Fixa sobre solo; Schletter; mount_id 17263; sem rastreamento.'],
       ['Orientação','Inclinação 25°; azimute 180° contado a partir do norte no sentido horário, voltado ao sul.'],
       ['Dimensões entre fileiras','Não informadas; não calcular sombreamento real sem esse dado.'],
       ['Albedo','0,20 assumido; sensibilidades 0,10 e 0,40 na transposição exploratória. A POA medida não recebe reflexão novamente.'],
       ['Sujidade','Base 0%; cenários 2% e 5% de atenuação adicional do módulo relativamente à leitura POA. Não são perdas medidas.'],
       ['Ventilação','Faiman genérico U₀=25 e U₁=6,84; sensibilidade de menor ventilação U₀=15 e U₁=3.']],
       [4.7,13],True,'Subseções 2.1 a 2.4. A inclinação é uma entrada do bloco de transposição, não uma nova medição solar.')
    head('3 Equipamentos e parâmetros da conversão fotovoltaica',True)
    tab('Módulos inversores e associação elétrica',['Elemento','Especificação e origem'],[
       ['Módulo no PVDAQ','Hyundai HiS-M310TI; 24 grupos de 120 módulos; 2.880 no total. Tipo no cadastro: mono-Si.'],
       ['Registro CEC selecionado',cfg['system']['module_cec_key']+'. Tecnologia do registro: '+str(module['Technology'])+'.'],
       ['Divergência de tecnologia','PVDAQ informa mono-Si e CEC informa Multi-c-Si. Preservam-se ambos os registros; a tecnologia instalada requer conferência com documentação do fabricante.'],
       ['Inversor no PVDAQ','ABB TRIO-27.6-TL-OUTD-S1B; tipo string; 24 unidades.'],
       ['Registro CEC de inversor',cfg['system']['inverter_cec_key']+'. A variante US 480 V selecionada não comprova a tensão ou a variante instalada.'],
       ['Topologia real de strings','modules_per_string e num_strings estão vazios. O inverter_id dos grupos de módulos está None. A associação completa não foi documentada.'],
       ['Equivalente didático','20 módulos em série × 6 strings em paralelo =120 módulos por grupo. Não é diagrama verificado da usina.']],
       [4.7,13],True)
    tab('Capacidades e relação entre DC e AC',['Grandeza','Valor','Unidade'],[
       ['Capacidade declarada no cadastro',893,'kWp'],['Soma de placa 2.880×310 W',892.8,'kWp'],
       ['Soma STC CEC 2.880×309,6 W',891.648,'kWp'],['Limite AC adotado 24×30 kW',720,'kW'],
       ['Nominal de comparação 24×27,6 kW',662.4,'kW'],['Relação DC/AC adotada 893/720',word.value(893/720,5),'1']],
       [10,3.8,3.9],note='O limite configurado da instalação não foi comprovado; 30 kW e 27,6 kW são hipóteses comparadas. Para Yf e PR usa-se P₀=893 kWp.')
    page()
    numeric=[(k,val) for k,val in module.items() if isinstance(val,(int,float))]
    tab('Todos os parâmetros numéricos do módulo CEC selecionado',['Parâmetro CEC','Valor','Unidade'],
       [[k,v(val),UNIT_MODULE.get(k,'indicador')] for k,val in numeric],[6.2,6.1,5.4],
       note='Arquivo config/modulo_cec.json. STC=309,6 W; I_mp×V_mp=8,6×36 V. N_s=60 é o valor do registro CEC, sem verificação do número físico de células. O diodo único usa alpha_sc, a_ref, I_L_ref, I_o_ref, R_s, R_sh_ref e Adjust.')
    tab('Campos descritivos restantes do módulo CEC',['Campo','Valor'],
       [[k,v(val)] for k,val in module.items() if not isinstance(val,(int,float))],[5.7,12],True,
       'Os campos Technology e BIPV são classificações do registro; Date e Version são da base CEC/SAM, não datas de medição ou fabricação deste sistema.')
    page()
    tab('Todos os parâmetros numéricos do inversor CEC selecionado',['Parâmetro CEC','Valor','Unidade'],
       [[k,v(val),UNIT_INV.get(k,'indicador')] for k,val in inv.items() if isinstance(val,(int,float))],[6.2,6.1,5.4],
       note='Inventário do registro CEC. O modelo principal do inversor usa a equação PVWatts, eficiência assumida e limite AC; não usa todos os coeficientes Sandia C0–C3. Não confundir Pdco deste registro com a potência STC do arranjo.')
    tab('Campos descritivos restantes do inversor CEC',['Campo','Valor'],
       [[k,v(val)] for k,val in inv.items() if not isinstance(val,(int,float))],[5.7,12],True,
       'Vac=480 é atributo do registro selecionado. Fabricante e modelo físicos provêm do PVDAQ; a variante elétrica instalada permanece por conferir.')
    page()
    rows=[]
    for k,val in cfg['assumptions'].items():
        if isinstance(val,bool) or k=='nasa_hour_label_offset_hours':continue
        rows.append([LABEL.get(k,k),v(val),UNIT_ASSUMPTIONS.get(k,'1')])
    tab('Parâmetros assumidos na cadeia principal',['Parâmetro','Valor','Unidade'],rows,[9.5,3.5,4.7],
       note='Hipóteses de modelagem, não medições independentes. As perdas DC são multiplicativas. A temperatura Faiman é usada como aproximação de célula. O deslocamento NASA de configuração é inativo porque NASA não entra na cadeia.')
    tab('Conversão e grandezas elétricas disponíveis',['Etapa ou variável','Utilização'],[
       ['POA e temperatura','POA efetiva e temperatura operacional alimentam calcparams_cec e o diodo único.'],
       ['Ponto de máxima potência','Tensão, corrente e potência do módulo são calculadas; não são ensaios I–V medidos.'],
       ['Potência DC','Soma ideal de 2.880 módulos, seguida das perdas assumidas.'],
       ['Potência AC','Conversão por PVWatts, com saturação no limite assumido.'],
       ['Canais dos inversores','Para cada uma das 24 unidades há corrente DC, tensão DC, corrente AC, tensão AC e potência AC.'],
       ['Potência DC por V×I','Diagnóstico derivado de canais observados; agregações e significado de cada canal limitam a comparação direta com DC simulada.'],
       ['Energia','Integral discreta das potências médias; Δt=0,25 h no passo comum.']],
       [5.2,12.5],True,'Subseções 3.1 a 3.6. Unidades dos canais elétricos: A, V e kW. Potências convertidas para W no cálculo.')
    page()
    tab('Inventário completo das 24 unidades inversoras',['Índice no cadastro','ID PVDAQ','Nome'],
       [[str(i),str(item['inverter_id']),item['name']] for i,item in enumerate(meta['Inverters'].values())],
       [5.7,5,7],True,'Todas têm fabricante ABB, modelo TRIO-27.6-TL-OUTD-S1B e quantidade 1. O CSV inversores contém também os campos vazios e o código temporal L, cuja semântica não foi inferida.')
    head('4 Avaliação qualidade e limites de validação',True)
    audit=summary['audit']
    tab('Disponibilidade e janelas de comparação',['Indicador','Valor','Unidade'],[
       ['Intervalos totais do trimestre',8736,'15 min'],['Medidor com dados',audit['meter_records'],'15 min'],
       ['Intervalos diurnos esperados',audit['expected_daylight_15m'],'15 min'],
       ['Pares diurnos disponíveis',audit['available_daylight_15m'],'15 min'],
       ['Pares após limites físicos',audit['paired_daylight_after_bounds'],'15 min'],
       ['Disponibilidade conjunta',word.value(100*4401/4421,4),'%'],['Intervalos diurnos sem par',20,'15 min']],
       [10.2,3.7,3.8],note='A cobertura mede disponibilidade, não certifica exatidão de cada sensor. Não houve preenchimento de lacunas ou exclusão por resíduo alto.')
    tab('Referências elétricas observadas',['Registro','Definição'],[
       ['Medidor AC','meter_revenue_grade_ac_output_meter_149578; cadastrado como revenue grade; fabricante Unknown; modelo e serial vazios.'],
       ['Soma dos inversores','Soma das potências AC das 24 unidades. É uma segunda referência local, com divergências próprias em relação ao medidor.'],
       ['Independência de entradas e saída','A geração medida não alimenta a linha física principal; serve para avaliar a saída AC calculada.'],
       ['Ajuste empírico separado','Fator '+v(summary['empirical_scale_fitted_only_april'])+' ajustado somente em abril. Maio e junho são avaliações temporais posteriores; não são uma instalação independente.']],
       [5,12.7],True)
    page()
    primary=summary['primary_validation'];units={'n':'pares','MBE':'kW','MAE':'kW','RMSE':'kW','nMBE_percent':'%','nRMSE_percent':'%','r':'1','R2':'1'}
    tab('Resultados da comparação elétrica principal',['Indicador','Valor','Unidade'],
       [[k,v(primary[k]),units[k]] for k in units],[7.5,6,4.2],
       note='Erro=simulado−medido. nMBE e nRMSE usam a média observada dos mesmos pares. R²=1−SSE/SST, calculado separadamente de r.')
    y=pd.read_csv(csv/'yields.csv').iloc[0]
    tab('Energia e desempenho com denominadores compatíveis',['Grandeza','Valor','Unidade'],[
       ['Energia observada integral do trimestre',word.value(summary['full_quarter']['E_AC_meter_full_quarter_kWh'],3),'kWh'],
       ['Energia observada na janela pareada',word.value(y.E_AC_observed_kWh,3),'kWh'],
       ['Energia AC simulada na janela pareada',word.value(y.E_AC_model_kWh,3),'kWh'],
       ['Irradiação POA na janela pareada',word.value(y.H_POA_kWh_m2,3),'kWh/m²'],
       ['Yf observado pareado',word.value(y.Yf_observed_kWh_kWp,3),'kWh/kWp'],
       ['PR observado pareado',word.value(y.PR_observed_paired*100,3),'%'],
       ['PR simulado pareado',word.value(y.PR_model*100,3),'%']],
       [10.2,4.2,3.3],note='Não combinar energia integral com POA parcial. HSP/PSH da janela=H_POA/(1 kW/m²); não é duração medida de brilho solar.')
    page()
    tab('Escopo da inversão e do BIPV',['Elemento','Dado e alcance'],[
       ['Total do trimestre',str(summary['reverse_POA']['total_quarter_hours'])+' horas.'],
       ['Recorte elegível da inversão',str(summary['reverse_POA']['candidate_hours'])+' horas; mesmos '+str(summary['reverse_POA']['converged_hours'])+' resultados convergentes.'],
       ['Métodos','Driesse 2024, Erbs–Driesse e Perez–Driesse; albedo de referência 0,20.'],
       ['Fechamento numérico','Resíduo máximo '+v(summary['reverse_POA']['max_numerical_closure_residual_Wm2'])+' W/m²; não comprova a exatidão das componentes.'],
       ['Geometrias comparadas','Horizontal; arranjo de referência a 25°/180°; fachadas sul, leste e oeste a 90°.'],
       ['Balanço semitransparente','Exemplo separado de 100 m²: energia refletida, transmitida, elétrica DC e calor para interior/exterior.'],
       ['Validação de fachada','Não há fachada medida, ensaio óptico ou SHGC certificado. Não se calculam iluminância nem demanda térmica de edifício.']],
       [5.1,12.6],True,'Subseção 4.5 é uma extensão didática da cadeia, não uma segunda usina observada.')
    tab('Todos os parâmetros numéricos do exemplo BIPV',['Parâmetro','Valor','Unidade'],
       [[LABEL.get(k,k),v(val),UNIT_BIPV.get(k,'1')] for k,val in cfg['bipv'].items() if isinstance(val,(int,float))],
       [9.5,3.5,4.7],note='Propriedades hipotéticas. Absortância=1−ρ−τsolar=0,63; a transmitância visível não é a transmitância solar. A eficiência elétrica depende da temperatura no exemplo.')
    head('Apêndice Dicionário completo dos 125 canais',True)
    word.paragraph(doc,'Cada linha identifica uma coluna original. A coluna measured_on é o carimbo temporal comum de cada CSV e não integra a contagem de 125 canais. O apêndice CSV inclui também escala, offset, agregação, amostras válidas e faltantes, mínimos, médias, máximos e zeros de cada canal no arquivo e no trimestre.')
    word.paragraph(doc,'No canal de potência AC do inversor 15, o cadastro usa inv_15_ac_power_iinv_149653 e o CSV usa inv_15_ac_power_inv_149653. A correspondência foi conferida pelo ID único 149653; ambos os nomes originais são preservados no dicionário, sem renomear o CSV.')
    def family(r):
        name=r.column
        if name.startswith('inv_'):
            bits=name.split('_');return bits[2].upper()+' '+{'power':'P','voltage':'V','current':'I'}[bits[3]]+' '+bits[1]
        return {'Irradiance POA':'POA','Temperature ambient':'T ar','Wind speed':'Vento','Wind direction':'Direção','AC power':'AC med.'}.get(r.common_name,r.common_name)
    tab('Todos os canais originais e suas unidades',['ID','Família','Unidade','Coluna do CSV'],
       [[str(int(r.metric_id)) if pd.notna(r.metric_id) else 'N/D',family(r),r.stored_unit if pd.notna(r.stored_unit) else 'N/D',r.column] for _,r in channels.iterrows()],
       [2.1,2.1,2.1,11.4],True,'Unidades e nomes são preservados do cadastro, inclusive kw. N/D indica campo não localizado no cadastro. O arquivo dicionario_completo_125_canais.csv contém as definições sem abreviar as famílias.')
    page()
    head('Apêndices digitais e referências de método')
    word.paragraph(doc,'Na pasta outputs/dossie_tecnico estão o JSON integral original, todos os campos de metadados em CSV, os inventários de arquivos, inversores, grupos de módulos, montagem, medidor e estação, o dicionário dos 125 canais e todos os campos da configuração do modelo. Os cinco arquivos originais continuam disponíveis em data/raw/pvdaq_2107_2024_originais.zip, conferidos por SHA-256.')
    for source in json.loads((ROOT/'data/raw/fontes_pvdaq.json').read_text()):
        word.paragraph(doc,source['file'],True)
        word.paragraph(doc,'SHA-256: '+source['sha256'])
    word.paragraph(doc,'As referências do projeto estão em references/ e o levantamento teórico por subseção está em bibliografia/outputs/. Os métodos centrais são SPA/NOAA, transposição Perez e Driesse, temperatura Faiman, diodo único CEC e inversor PVWatts. A matriz de cobertura relaciona cada método à subseção e ao responsável.')
    dest=out/'Dossie_Tecnico_por_Secao.docx';word.save_document(doc,dest)
    manifest={'original_files_verified':5,'measured_channels':125,'metadata_leaf_fields':len(flatten(meta)),
       'module_parameter_fields':len(module),'inverter_parameter_fields':len(inv),
       'matched_metric_ids':125,'metadata_name_aliases':1,
       'chapter_sections':4,'numeric_tables':nums['Tabela'],'text_frames':nums['Quadro'],
       'source_year':2026,'observation_year':2024,'status':'gerado',
       'document':str(dest.relative_to(ROOT)),'known_discrepancies':['PVDAQ mono-Si versus CEC Multi-c-Si',
       'datas de início do sistema e dos grupos de módulos diferentes','metadados históricos terminam em 2023 mas CSV disponíveis em 2024',
       'variante e limite AC do inversor sem comprovação; topologia de strings não informada',
       'nome do canal de potência AC do inversor 15 difere no cadastro; correspondência por ID único 149653']}
    (out/'inventario_do_dossie.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
    print(json.dumps(manifest,ensure_ascii=False,indent=2));print(dest,flush=True)
    return dest

if __name__=='__main__':main()
