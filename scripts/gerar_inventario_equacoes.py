"""Curated equation inventory with native Word math and exclusive insertion IDs.

N records transcribe the mathematical relations of the working notebook at
BASE. C records expose complementary relations already used by its code.
No physical model, parameter, data or numerical output is rewritten.
"""
from pathlib import Path
from collections import Counter
from copy import deepcopy
import csv
import hashlib
import json
import re
import subprocess
import tempfile
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from metadados_equacoes import enrich

ROOT=Path(__file__).resolve().parents[1]
BASE='17ffe0eeb7d70876eabf222ec30fb828ee28d3c8'
REPOSITORY='https://github.com/samuelasroma-byte/SPGER0046_Energias_Renovaveis'
MNS='http://schemas.openxmlformats.org/officeDocument/2006/math'
FOLDERS=['01_Recurso_Solar','02_Condicoes_Ambientais','03_Modelagem_Fotovoltaica','04_Desempenho_Perdas_Validacao']
TITLES=['Fundamentos e caracterização do recurso solar',
        'Condições ambientais e de entorno da conversão fotovoltaica',
        'Modelagem da conversão fotovoltaica',
        'Avaliação de desempenho perdas e validação da modelagem']
OWNERS=['Samuel','Vannessa e Thais','Gabriela e Giulia','Júlia Ribeiro e Lucas Simão']


def inventory():
    source=json.loads(subprocess.check_output(['git','-C',str(ROOT),'show',BASE+':docs/secoes_capitulo.json'],text=True))
    by_sub={r['id']:r for r in source}
    rows=[];sequence=Counter()

    def add(sub,kind,title,formula,function,symbols,locator,refs=None,notes=''):
        section=int(sub[0]);sequence[section]+=1
        record=dict(id=f'E{section}.{sequence[section]:02d}',section=section,subsection=sub,
                    title=title,latex=formula,function=function,symbols=symbols,
                    origin=kind,source_locator=locator,
                    refs=refs if refs is not None else by_sub[sub]['refs'],notes=notes)
        rows.append(record);return record

    def n(sub,index,title,function,symbols,locator,notes='',formula=None):
        fragments=re.split(r',\\(?:qquad|quad)\s*',by_sub[sub]['equation'].strip('$').strip().rstrip('.'))
        original=fragments[index]
        record=add(sub,'N',title,formula or original,function,symbols,locator,notes=notes)
        record['notebook_fragment']=original
        record['notebook_subsections']=[sub]
        return record

    def c(sub,title,formula,function,symbols,locator,notes='',refs=None):
        return add(sub,'C',title,formula,function,symbols,locator,refs,notes)

    n('1.1',0,'Ângulo zenital','Relacionar a posição do Sol à latitude e ao tempo solar.',
      'φ latitude; δ declinação; ω ângulo horário; θz zênite.','src/projeto_fv.py::geometry',
      'Relação geométrica de fundamentação; o cálculo utiliza SPA e o meio de cada intervalo.')
    n('1.1',1,'Irradiância extraterrestre horizontal','Projetar a irradiância extraterrestre normal na horizontal.',
      'G0n e G0h em W/m².','src/projeto_fv.py::geometry')
    n('1.2',0,'Componentes na horizontal','Distinguir irradiâncias global horizontal direta normal e difusa horizontal.',
      'GHI DNI e DHI em W/m².','docs/secoes_capitulo.json::1.2',
      'Identidade de fundamentação; GHI DNI e DHI locais são estimativas na exploração inversa.')
    n('1.2',1,'Irradiação acumulada','Integrar irradiância média no plano e período definidos.',
      'G em W/m²; Δt em h; H em Wh/m²; dividir por 1000 para kWh/m².',
      'src/refinamento_fv.py::analytics')
    n('1.3',0,'Variabilidade relativa da irradiância','Ordenar os dias pelo passo absoluto médio da POA relativo à sua média.',
      'Gi POA dos pares válidos ordenados do dia; resultado adimensional.',
      'src/refinamento_fv.py::analytics',
      'O código aplica diff aos pares disponíveis do dia; intervalos entre pares podem conter lacunas.')
    n('1.4',0,'Cobertura dos dados','Informar a proporção de intervalos válidos frente aos esperados.',
      'Npares e Nesperados na mesma janela; C em %.','src/refinamento_fv.py::analytics')

    c('2.1','Média amostral',r'\bar{x}=\frac{1}{n}\sum_{i=1}^{n}x_i',
      'Resumir a posição das distribuições ambientais e das variáveis tabuladas.',
      'xi variável analisada; n pares; média na unidade de x.','src/refinamento_fv.py::analytics',refs=[])
    c('2.1','Desvio padrão amostral',r's_x=\sqrt{\frac{\sum_{i=1}^{n}(x_i-\bar{x})^2}{n-1}}',
      'Quantificar a dispersão com a convenção amostral usada no projeto.',
      'n maior que 1; sx na unidade de x.','src/refinamento_fv.py::analytics',
      'O código usa std com ddof igual a 1. A equação de Faiman fica em E3.03 e deve ser apenas citada em 2.1.',refs=[])
    c('2.1','Amplitude interquartil',r'IQR=Q_3-Q_1',
      'Resumir a dispersão central entre os percentis 25 e 75.',
      'Q1 e Q3 quartis da mesma variável e janela.','src/refinamento_fv.py::analytics',refs=[])
    n('2.2',0,'Reflexão isotrópica do solo','Estimar a parcela refletida pelo solo que incide no plano inclinado.',
      'ρg albedo; β inclinação; irradiâncias em W/m².','src/projeto_fv.py::transpose',
      'Aplicada aos cenários de transposição; não somar novamente à POA medida da cadeia principal.')
    n('2.3',0,'Irradiância efetiva','Aplicar a hipótese de atenuação adicional à leitura POA.',
      's fração assumida entre 0 e 1; Gef e GPOA em W/m².','src/projeto_fv.py::simulate',
      'A atenuação uniforme não representa sombreamento parcial com mismatch ou diodos de bypass.')
    n('2.4',0,'Ângulo de incidência','Relacionar a geometria do Sol à orientação do plano.',
      'β inclinação; γs e γa azimutes solar e do arranjo; θi incidência.',
      'src/projeto_fv.py::geometry')
    n('2.4',1,'Projeção geométrica de sombra','Ilustrar a projeção horizontal de uma altura sob o Sol.',
      'h altura; αs altura solar positiva; ℓ comprimento na unidade de h.',
      'docs/secoes_capitulo.json::2.4',
      'Exemplo plano didático; não calcula o sombreamento real das fileiras sem suas dimensões.')

    n('3.1',0,'Transposição para o plano do arranjo','Combinar radiação direta difusa do céu e refletida pelo solo.',
      'Gcéu e Gsolo no plano; todos os termos em W/m².','src/projeto_fv.py::transpose',
      'Cenários usam Perez–Driesse; a cadeia principal recebe POA medida diretamente.')
    c('3.1','Resíduo de fechamento da POA',r'\varepsilon_{\mathrm{POA}}=G_{\mathrm{POA,recon}}-G_{\mathrm{POA,med}}',
      'Verificar o fechamento numérico da transposição inversa.',
      'POA reconstruída e medida no mesmo intervalo; resíduo em W/m².',
      'src/projeto_fv.py::reverse_components',
      'O fechamento não constitui validação independente das componentes reconstruídas.')
    thermal=n('3.2',0,'Temperatura operacional de Faiman','Estimar a temperatura operacional a partir da irradiância e da meteorologia.',
      'Ta em °C; v em m/s; U0 e U1 coeficientes térmicos.',
      'src/projeto_fv.py::simulate',
      'Temperatura operacional usada como aproximação da célula; sem medição térmica independente.',
      formula=r'T_{\mathrm{op}}=T_a+\frac{G_{\mathrm{ef}}}{U_0+U_1v}')
    thermal['notebook_subsections']=['2.1','3.2']
    n('3.3',0,'Modelo de diodo único','Relacionar corrente e tensão do módulo nas condições operacionais.',
      'IL e I0 correntes; Rs e Rsh resistências; a fator em volts.',
      'src/projeto_fv.py::module_mpp',
      'Parâmetros corrigidos pelo modelo CEC; curvas calculadas não são varreduras experimentais.',
      formula=r'I=I_L-I_0\left[\exp\left(\frac{V+IR_s}{a}\right)-1\right]-\frac{V+IR_s}{R_{sh}}')
    n('3.3',1,'Potência elétrica','Obter potência e construir curvas P–V e o ponto de máxima potência.',
      'V em V; I em A; P em W.','src/projeto_fv.py::module_curves')
    c('3.3','Comparação CC com PVWatts',
      r'P_{\mathrm{DC,PVW}}=P_{\mathrm{STC}}\frac{G_{\mathrm{ef}}}{1000}\left[1+\gamma(T_{\mathrm{op}}-25)\right]\prod_j(1-\ell_j)',
      'Registrar a aproximação linear CC exportada como comparação auxiliar.',
      'Gef em W/m²; γ em °C⁻¹; PSTC potência total dos módulos em W.',
      'src/projeto_fv.py::simulate',
      'Comparação auxiliar; a previsão principal usa o modelo CEC de diodo único.',refs=['R11','R09'])
    n('3.4',0,'Tensão em série','Calcular a tensão do equivalente de módulos em série.',
      'Ns número em série; Vm tensão do módulo.','src/refinamento_fv.py::analytics')
    n('3.4',1,'Corrente em paralelo','Calcular a corrente do equivalente de strings em paralelo.',
      'Np strings em paralelo; Im corrente do módulo.','src/refinamento_fv.py::analytics')
    n('3.4',2,'Potência do arranjo ideal','Escalar a potência do módulo para a associação ideal.',
      'Pm potência do módulo; Pa em W.','src/refinamento_fv.py::analytics',
      'O exemplo 20 em série e 6 em paralelo não comprova o cabeamento real; a simulação total usa 2880 módulos.')
    n('3.4',3,'Perdas CC multiplicativas','Aplicar perdas assumidas de cabos mismatch e conexões.',
      'ℓj frações de perda; Pa potência bruta; PDC potência líquida.',
      'src/projeto_fv.py::simulate')
    n('3.5',0,'Potência CA limitada','Limitar a conversão pela capacidade CA assumida.',
      'η eficiência; PACmax limite; potências em W.','src/projeto_fv.py::inverter_pvwatts_stages',
      'No código o piso zero é aplicado antes do limite superior; usar também E3.15.')
    n('3.5',1,'Referência CC do inversor','Converter a capacidade CA em potência CC de referência da curva.',
      'ηnom eficiência nominal; Pdc0 não é a potência STC dos módulos.',
      'src/projeto_fv.py::inverter_pvwatts_stages')
    c('3.5','Carregamento relativo',r'\zeta=\frac{P_{\mathrm{DC}}}{P_{\mathrm{dc0,inv}}}',
      'Definir a variável independente da curva de eficiência.',
      'ζ adimensional; mesma unidade nas duas potências.',
      'src/projeto_fv.py::inverter_pvwatts_stages')
    c('3.5','Eficiência PVWatts',r'\eta=\frac{\eta_{\mathrm{nom}}}{\eta_{\mathrm{ref}}}\left(-0.0162\zeta-\frac{0.0059}{\zeta}+0.9858\right)',
      'Calcular a eficiência em função do carregamento do inversor.',
      'ζ positivo; ηnom igual a 0,98; ηref igual a 0,9637.',
      'src/projeto_fv.py::inverter_pvwatts_stages',
      'Em potência CC nula o código define saída CA nula sem avaliar a divisão por ζ.')
    c('3.5','Potência antes do clipping',r'P_{\mathrm{pre}}=\max(\eta P_{\mathrm{DC}},0)',
      'Impor o piso físico não negativo antes da limitação superior.',
      'Ppre em W; PDC igual a zero produz Ppre igual a zero.',
      'src/projeto_fv.py::inverter_pvwatts_stages')
    c('3.5','Potência suprimida pelo clipping',r'P_{\mathrm{clip}}=P_{\mathrm{pre}}-P_{\mathrm{AC}}',
      'Quantificar a parcela limitada pela capacidade do inversor.',
      'Pclip em W; não confundir com energia perdida.',
      'src/projeto_fv.py::inverter_pvwatts_stages')
    c('3.5','Potência CA final',r'P_{\mathrm{AC,final}}=(1-\ell_{\mathrm{AC}})P_{\mathrm{AC}}',
      'Aplicar a perda CA assumida após a conversão e o clipping.',
      'ℓAC fração de perda; PACfinal em W.','src/projeto_fv.py::simulate')
    n('3.6',0,'Energia elétrica','Integrar a potência média nos intervalos selecionados.',
      'P em kW; Δt igual a 0,25 h; E em kWh. Se P estiver em W dividir por 1000.',
      'src/projeto_fv.py::daily_paired_energy',
      'Energia pareada e energia integral do mês são totais diferentes e devem ser identificados.')

    n('4.1',0,'Variação de energia entre estágios','Quantificar a diferença calculada entre dois estágios consecutivos.',
      'Ej energias na mesma máscara temporal em kWh.','src/projeto_fv.py::performance')
    n('4.1',1,'Fechamento da cadeia de energia','Conferir que a soma das variações recupera a diferença total.',
      'Mesmos estágios e intervalos; podem ocorrer ganhos sob frio.',
      'src/projeto_fv.py::performance',
      'Decomposição do modelo não identifica automaticamente perdas reais ou suas causas.')
    c('4.1','Potência linear de referência',r'P_{\mathrm{lin}}=P_0\frac{G_{\mathrm{POA}}}{G_{\mathrm{ref}}}',
      'Definir a referência linear STC que inicia a cadeia energética.',
      'P0 capacidade declarada; Gref igual a 1000 W/m²; potências na mesma unidade.',
      'src/projeto_fv.py::simulate',
      'Na cadeia principal P0 declarada é 893 kWp; não substituir pela soma CEC sem identificar a mudança.')
    n('4.2',0,'Produtividade de referência','Expressar o recurso como horas equivalentes no plano e período definidos.',
      'HPOA em kWh/m²; Gref igual a 1 kW/m²; Yr em h.',
      'src/projeto_fv.py::performance',
      'HSP e PSH designam o mesmo conceito nesta janela; não são duração de brilho solar.')
    n('4.2',1,'Produtividade do arranjo','Normalizar a energia CC pela capacidade declarada.',
      'EDC em kWh; P0 em kWp; Ya em kWh/kWp ou horas equivalentes.',
      'src/projeto_fv.py::performance','A energia CC e Ya são modeladas neste estudo.')
    n('4.2',2,'Produtividade final','Normalizar a energia CA observada ou modelada pela capacidade.',
      'EAC em kWh; P0 em kWp; Yf em kWh/kWp.','src/projeto_fv.py::performance')
    n('4.3',0,'Razão de desempenho','Relacionar a produtividade final ao recurso da mesma janela.',
      'PR adimensional; multiplicar por 100 quando apresentado em %.','src/projeto_fv.py::performance')
    n('4.3',1,'Perda de captura normalizada','Comparar o recurso e a produtividade CC do modelo.',
      'Lc em horas equivalentes.','src/projeto_fv.py::performance')
    n('4.3',2,'Perda de sistema normalizada','Comparar as produtividades CC e CA do modelo.',
      'Ls em horas equivalentes.','src/projeto_fv.py::performance')
    c('4.4','Resíduo de potência',r'e_i=p_i-y_i',
      'Fixar o sinal do erro usado nas tabelas e gráficos.',
      'pi potência simulada; yi observada; ambas em kW.',
      'src/projeto_fv.py::metrics','Valores positivos indicam superestimação.')
    n('4.4',0,'Erro médio de viés','Quantificar a tendência média de superestimação ou subestimação.',
      'n pares válidos; MBE em kW.','src/projeto_fv.py::metrics',
      formula=r'MBE=\frac{1}{n}\sum_{i=1}^{n}e_i')
    n('4.4',1,'Erro absoluto médio','Quantificar a magnitude média do erro sem cancelamento de sinais.',
      'MAE em kW.','src/projeto_fv.py::metrics',formula=r'MAE=\frac{1}{n}\sum_{i=1}^{n}|e_i|')
    n('4.4',2,'Raiz do erro quadrático médio','Resumir a magnitude quadrática dos mesmos resíduos.',
      'RMSE em kW.','src/projeto_fv.py::metrics',formula=r'RMSE=\sqrt{\frac{1}{n}\sum_{i=1}^{n}e_i^2}')
    c('4.4','Viés normalizado',r'nMBE=100\frac{MBE}{\bar y}',
      'Expressar o viés em relação à média observada dos pares.',
      'nMBE em %; média observada não nula.','src/projeto_fv.py::metrics')
    c('4.4','Erro quadrático normalizado',r'nRMSE=100\frac{RMSE}{\bar y}',
      'Expressar o RMSE em relação à média observada dos pares.',
      'nRMSE em %; não usar potência nominal como denominador neste projeto.',
      'src/projeto_fv.py::metrics')
    c('4.4','Correlação de Pearson',r'r=\frac{\sum_i(y_i-\bar y)(p_i-\bar p)}{\sqrt{\sum_i(y_i-\bar y)^2\sum_i(p_i-\bar p)^2}}',
      'Quantificar a associação linear entre observações e simulações.',
      'r adimensional; requer variância positiva nas duas séries.',
      'src/projeto_fv.py::metrics','Correlação elevada não prova exatidão.')
    n('4.4',3,'Coeficiente de determinação','Comparar o erro quadrático à dispersão das observações.',
      'R² adimensional; denominador positivo.','src/projeto_fv.py::metrics',
      'O projeto calcula 1 menos SSE/SST; não substitui R² pelo quadrado de r.')
    c('4.4','Fator empírico de abril',r'k=\frac{\sum_{i\in\mathrm{abril}}p_i y_i}{\sum_{i\in\mathrm{abril}}p_i^2}',
      'Ajustar um fator multiplicativo secundário apenas com os pares de abril.',
      'k adimensional; pi e yi na mesma unidade.','src/projeto_fv.py::run_models',
      'Ajuste sem intercepto; maio é verificação e junho avaliação temporal separada.',refs=[])
    c('4.4','Potência com ajuste empírico',r'p_i^{\mathrm{emp}}=k p_i',
      'Aplicar o fator de abril e reportar os resultados separadamente do modelo físico.',
      'Potências na mesma unidade; k fixado no mês de ajuste.',
      'src/projeto_fv.py::run_models',refs=[])
    c('4.5','Absortância solar',r'\alpha=1-\rho_s-\tau_s',
      'Fechar a partição óptica do exemplo semitransparente.',
      'α absortância; ρs refletância; τs transmitância solar; frações.',
      'src/projeto_fv.py::bipv_scenario',
      'As propriedades são hipóteses do exemplo de 100 m² e não do módulo real.')
    c('4.5','Potência solar incidente',r'P_{\mathrm{inc}}=GA',
      'Converter a irradiância do plano em potência incidente na área.',
      'G em W/m²; A em m²; Pinc em W.','src/projeto_fv.py::bipv_scenario')
    c('4.5','Potência solar refletida',r'P_{\mathrm{ref}}=\rho_s GA',
      'Quantificar a fração refletida da potência incidente.',
      'Pref em W.','src/projeto_fv.py::bipv_scenario')
    c('4.5','Potência solar transmitida',r'P_{\mathrm{trans}}=\tau_s GA',
      'Quantificar a parcela solar que atravessa o elemento.',
      'Ptrans em W; τs solar difere da transmitância visível.',
      'src/projeto_fv.py::bipv_scenario')
    c('4.5','Eficiência elétrica do exemplo BIPV',
      r'\eta=\min\{\max[\eta_{\mathrm{ref}}(1+\gamma(T_{\mathrm{op}}-25)),0],\alpha\}',
      'Corrigir a eficiência assumida pela temperatura e limitar a conversão à absorção.',
      'ηref eficiência a 25 °C; γ em °C⁻¹.','src/projeto_fv.py::bipv_scenario',
      'Esta aproximação pertence ao balanço semitransparente; os cenários de produtividade usam a cadeia CEC.')
    n('4.5',0,'Potência elétrica do elemento BIPV','Separar a parcela convertida em eletricidade no balanço solar.',
      'η eficiência instantânea; Pe em W.','src/projeto_fv.py::bipv_scenario')
    n('4.5',1,'Calor absorvido após conversão elétrica','Calcular a parcela absorvida que permanece térmica.',
      'Qabs representa potência térmica em W apesar do símbolo Q.',
      'src/projeto_fv.py::bipv_scenario')
    c('4.5','Calor absorvido dirigido ao exterior',r'Q_{\mathrm{out}}=(1-f_{\mathrm{in}})Q_{\mathrm{abs}}',
      'Separar a fração do calor absorvido dirigida ao exterior.',
      'fin fração dirigida ao interior entre 0 e 1; Qout em W.',
      'src/projeto_fv.py::bipv_scenario')
    n('4.5',2,'Ganho solar para o interior','Somar transmissão solar e fração térmica dirigida ao interior.',
      'Qin é potência em W; fin fração térmica interior.',
      'src/projeto_fv.py::bipv_scenario',
      'Balanço solar simplificado; não inclui todas as trocas térmicas do edifício.')
    n('4.5',3,'SHGC equivalente simplificado','Normalizar o ganho solar interior pela potência incidente.',
      'SHGCeq adimensional; não equivale a certificação de produto.',
      'src/projeto_fv.py::bipv_scenario')
    c('4.5','Conservação do balanço solar',
      r'P_{\mathrm{inc}}=P_{\mathrm{ref}}+P_{\mathrm{trans}}+P_e+f_{\mathrm{in}}Q_{\mathrm{abs}}+Q_{\mathrm{out}}',
      'Conferir a conservação das parcelas ópticas elétricas e térmicas.',
      'Todos os termos em W; o código testa o resíduo de fechamento.',
      'src/projeto_fv.py::bipv_scenario')
    assert thermal['id']=='E3.03'
    assert len(rows)==61 and Counter(r['origin'] for r in rows)=={'N':37,'C':24}
    # Stable cross-reference to the nonnegative floor in the inverter model.
    floor=next(r['id'] for r in rows if r['title']=='Potência antes do clipping')
    for r in rows:
        if r['title']=='Potência CA limitada':r['notes']=r['notes'].replace('E3.15',floor)
    enrich(rows)
    return dict(version='2026-10-19.2',source_commit=BASE,
                annotation_base_commit='1c1c38375ae106b7c6317958ebebbd6ad0e94e9a',
                title='Inventário de equações do capítulo de recurso solar e conversão fotovoltaica',
                count_rule='Uma relação numerada por linha. N indica fórmula escrita nas células de método; C explicita relação complementar do código. A classe de aplicação informa o uso efetivo e é independente de N/C. Constantes, parâmetros, conversões, máscaras e algoritmos internos das bibliotecas não recebem numeração editorial.',
                application_scope='A previsão principal do modelo A usa POA medida, atenuação assumida, Faiman, CEC/diodo único, perdas CC, inversor PVWatts e perdas CA. Integração, estatísticas e indicadores avaliam os dados e as previsões. Geometria, transposição inversa, curvas ilustrativas, comparação CC PVWatts, ajuste empírico e balanço BIPV têm funções próprias identificadas por equação.',
                notation_conventions=[
                    'Símbolos e unidades são definidos por equação, incluindo o resultado e os índices. A unidade 1 indica grandeza adimensional; % indica apresentação percentual.',
                    'Ângulos são declarados em graus. Rotinas trigonométricas que recebem radianos exigem conversão. Azimutes partem do norte no sentido horário.',
                    'Em E1.04, G em W/m² e Δt em h produzem H em Wh/m²; o código divide por 1000 para kWh/m². Em E3.18, P em kW produz E em kWh; o código converte suas séries W por /1000.',
                    'Em E3.06, 1000 representa Gref=1000 W/m² e 25 representa 25 °C. Em E4.24, 25 também representa 25 °C. Os coeficientes da curva E3.14 são adimensionais.',
                    'U0 tem unidade W/(m²·K) e U1 tem W·s/(m³·K), de modo que Gef/(U0+U1v) é diferença de temperatura em K, numericamente igual a °C.',
                    'P0 é 893000 W em E4.03 e 893 kWp nos indicadores. Gref é 1000 W/m² na potência e 1 kW/m² na produtividade de referência.',
                    'ηref=0,9637 é referência da curva do inversor em E3.14; ηref=0,10 é hipótese de eficiência BIPV em E4.24. γ vem do módulo CEC em E3.06 e é assumido −0,0044 °C⁻¹ em E4.24.',
                    'Pa em E3.09 refere-se ao exemplo de 120 módulos; Pa em E3.10 é a potência bruta dos 2880 módulos do sistema. A associação 20×6 não comprova o cabeamento real.',
                    'Qabs, Qin e Qout no balanço BIPV designam potências em W. A integração dessas séries fornece energia; Q não representa joules nessas equações.',
                    'Nas métricas locais de potência, p, y e e estão em kW; metrics é genérica e conserva a unidade das séries recebidas. O ajuste k é calculado em W no código e é invariável sob conversão comum para kW.',
                ],
                equations=rows,
                cross_references=[dict(subsection='2.1',equation_id='E3.03',reason='A relação de Faiman é inserida em 3.2 e apenas citada em 2.1.')],
                parameters=dict(reference_irradiance_W_m2=1000,interval_hours=0.25,
                                declared_capacity_kWp=893,eta_nom=0.98,eta_ref_inverter=0.9637,
                                eta_ref_bipv=0.10,bipv_gamma_per_C=-0.0044))


def counts(rows):
    return Counter(r['origin'] for r in rows)


def math_elements(rows,folder):
    markdown=folder/'matematica.md';target=folder/'matematica.docx'
    markdown.write_text('\n\n'.join('$$'+r['latex']+'$$' for r in rows))
    p=subprocess.run(['pandoc',str(markdown),'-o',str(target)],check=True,capture_output=True,text=True)
    assert not p.stderr.strip(),p.stderr
    with zipfile.ZipFile(target) as archive:
        xml=etree.fromstring(archive.read('word/document.xml'))
    equations=xml.findall('.//{'+MNS+'}oMath')
    assert len(equations)==len(rows)
    # Use the native bar construct so the mean symbol stays a horizontal bar
    # in both Word and LibreOffice, rather than a font-dependent accent glyph.
    for equation in equations:
        for accent in list(equation.findall('.//{'+MNS+'}acc')):
            character=accent.find('{'+MNS+'}accPr/{'+MNS+'}chr')
            if character is not None and character.get('{'+MNS+'}val')=='‾':
                bar=OxmlElement('m:bar');properties=OxmlElement('m:barPr')
                position=OxmlElement('m:pos');position.set(qn('m:val'),'top');properties.append(position)
                bar.append(properties);bar.append(deepcopy(accent.find('{'+MNS+'}e')))
                accent.getparent().replace(accent,bar)
        # LibreOffice can replace the closing delimiter of nested OMML groups
        # by a parenthesis. Explicit native math runs preserve square brackets,
        # braces and absolute-value bars in both the DOCX and its QA render.
        for delimiter in reversed(equation.findall('.//{'+MNS+'}d')):
            properties=delimiter.find('{'+MNS+'}dPr')
            begin=properties.find('{'+MNS+'}begChr') if properties is not None else None
            end=properties.find('{'+MNS+'}endChr') if properties is not None else None
            if begin is None or end is None:continue
            left,right=begin.get(qn('m:val')),end.get(qn('m:val'))
            if (left,right) not in [('[',']'),('{','}'),('|','|')]:continue
            expressions=delimiter.findall('{'+MNS+'}e');assert len(expressions)==1
            parent=delimiter.getparent();position=parent.index(delimiter)
            def literal(character):
                run=OxmlElement('m:r');props=OxmlElement('m:rPr')
                style=OxmlElement('m:sty');style.set(qn('m:val'),'p');props.append(style);run.append(props)
                text=OxmlElement('m:t');text.text=character;run.append(text);return run
            replacement=[literal(left),*[deepcopy(e) for e in expressions[0]],literal(right)]
            for offset,element in enumerate(replacement):parent.insert(position+offset,element)
            parent.remove(delimiter)
    return {r['id']:deepcopy(e) for r,e in zip(rows,equations)}


def border_table(table):
    table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
    props=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'6');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    props.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for edge,value in [('top',80),('left',120),('bottom',80),('right',120)]:
        e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');margins.append(e)
    props.append(margins)
    header=table.rows[0]._tr.get_or_add_trPr();header.append(OxmlElement('w:tblHeader'))
    for row_index,row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index==0:
                shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E8EEF4');cell._tc.get_or_add_tcPr().append(shd)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0)
                p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1
                if row_index==0:p.paragraph_format.keep_with_next=True
                for run in p.runs:
                    run.font.name='Times New Roman';run.font.size=Pt(12);run.font.color.rgb=RGBColor(0,0,0)
                    if row_index==0:run.bold=True


def document_base(title):
    doc=Document()
    section=doc.sections[0];section.orientation=WD_ORIENT.LANDSCAPE
    section.page_width=Cm(29.7);section.page_height=Cm(21)
    section.top_margin=Cm(2);section.bottom_margin=Cm(2)
    section.left_margin=Cm(2);section.right_margin=Cm(2)
    for style_name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        style=doc.styles[style_name];style.font.name='Times New Roman';style.font.size=Pt(12)
        style.font.color.rgb=RGBColor(0,0,0)
        fonts=style.element.find('.//'+qn('w:rFonts'))
        if fonts is not None:
            for name in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
                fonts.attrib.pop(qn('w:'+name),None)
        style.paragraph_format.space_after=Pt(5)
        for border in style.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].font.bold=True
    doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 1'].font.bold=True
    doc.styles['Heading 2'].font.size=Pt(12);doc.styles['Heading 2'].font.bold=True
    doc.add_paragraph(title,'Title')
    doc.core_properties.author='Equipe SPGER0046 PPGER UFPB'
    doc.core_properties.title=title
    doc.core_properties.subject='Inventário editorial de equações por seção e subseção'
    p=section.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run('SPGER0046   ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
    for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(10)
    return doc


def summary_table(doc,rows,sections):
    table=doc.add_table(rows=1,cols=4)
    for cell,label in zip(table.rows[0].cells,['Seção','Notebook N','Código C','Total de relações']):cell.text=label
    for section in sections:
        selected=[r for r in rows if r['section']==section];cc=counts(selected)
        cells=table.add_row().cells
        for cell,value in zip(cells,[str(section)+' '+TITLES[section-1],str(cc['N']),str(cc['C']),str(len(selected))]):cell.text=value
    if len(sections)>1:
        cc=counts(rows)
        for cell,value in zip(table.add_row().cells,['Total',str(cc['N']),str(cc['C']),str(len(rows))]):cell.text=value
    for column,width in zip(table.columns,[16.1,3,3,3.6]):column.width=Cm(width)
    border_table(table)
    for row in table.rows:
        for cell in row.cells[1:]:
            for p in cell.paragraphs:p.alignment=WD_ALIGN_PARAGRAPH.CENTER


def equation_table(doc,selected,maths):
    table=doc.add_table(rows=1,cols=4)
    for cell,label in zip(table.rows[0].cells,['ID e inserção','Equação','Função e símbolos com unidades','Registro e aplicação']):cell.text=label
    widths=[2.1,10.5,8.8,4.3]
    for column,width in zip(table.columns,widths):column.width=Cm(width)
    for r in selected:
        cells=table.add_row().cells
        for cell,width in zip(cells,widths):cell.width=Cm(width)
        cells[0].text=r['id']+'\n'+r['subsection']
        p=cells[1].paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        equation=deepcopy(maths[r['id']])
        # Word equations remain native searchable editable OMML, not images.
        for run in equation.findall('.//{'+MNS+'}r'):
            properties=run.find(qn('w:rPr'))
            if properties is None:properties=OxmlElement('w:rPr');run.insert(0,properties)
            fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Cambria Math');fonts.set(qn('w:hAnsi'),'Cambria Math');properties.append(fonts)
            size=OxmlElement('w:sz');size.set(qn('w:val'),'26');properties.append(size)
        p._p.append(equation)
        p=cells[2].paragraphs[0];p.add_run(r['title']+'. ').bold=True
        p.add_run(r['function'])
        cells[2].add_paragraph(r['symbols'])
        cells[3].text=('N Fórmula do notebook' if r['origin']=='N' else 'C Relação do código')+'\n'+r['application']
        if r['refs']:cells[3].add_paragraph(', '.join(r['refs']))
    border_table(table)
    for row in table.rows[1:]:
        for p in row.cells[0].paragraphs:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return table


def make_doc(data,maths,path,section=None):
    rows=[r for r in data['equations'] if section is None or r['section']==section]
    selected_sections=list(range(1,5)) if section is None else [section]
    title='Inventário geral de equações do capítulo' if section is None else 'Quadro de equações da seção '+str(section)
    doc=document_base(title)
    doc.add_paragraph('Recurso solar e modelagem da conversão fotovoltaica   PPGER UFPB   19 de outubro de 2026')
    if section:doc.add_paragraph(TITLES[section-1]+'. Responsáveis: '+OWNERS[section-1]+'.')
    cc=counts(rows)
    doc.add_paragraph(f'Este inventário organiza {len(rows)} relações numeradas: {cc["N"]} escritas nas células de método do notebook e {cc["C"]} complementares explicitadas a partir do código. '
                      'Usar a coluna de inserção para distribuir a redação e manter os IDs nas referências cruzadas.')
    doc.add_paragraph('A coluna de aplicação distingue conversão principal, tratamento, avaliação, fundamentação, ilustrações e cenários. '
                      'N e C indicam onde a relação foi registrada e não separam modelos físicos de cenários. Cada símbolo recebe descrição e unidade na própria linha; 1 indica adimensional.')
    doc.add_paragraph('O modelo A converte POA medida e meteorologia local por atenuação assumida, Faiman, CEC de diodo único, perdas CC, inversor PVWatts e perdas CA. '
                      'A transposição exploratória, a aproximação CC PVWatts, o ajuste empírico e o balanço BIPV são identificados em suas respectivas aplicações.')
    summary_table(doc,rows,selected_sections)
    doc.add_paragraph('O modelo de Faiman corresponde a E3.03, na subseção 3.2; a subseção 2.1 utiliza essa referência. '
                      '')
    for number in selected_sections:
        heading=doc.add_heading('Seção '+str(number)+' '+TITLES[number-1],level=1)
        heading.paragraph_format.page_break_before=False
        selected=[r for r in rows if r['section']==number]
        sub_counts=Counter(r['subsection'] for r in selected)
        distribution=doc.add_paragraph('Distribuição por subseção: '+'; '.join(f'{sub} com {sub_counts[sub]} '+('relação' if sub_counts[sub]==1 else 'relações') for sub in sorted(sub_counts))+'.')
        distribution.paragraph_format.keep_with_next=True
        if number==2:doc.add_paragraph('Em 2.1 as relações de apoio descrevem as estatísticas tabuladas. O modelo térmico deve ser citado em E3.03 sem repetir sua inserção.')
        equation_table(doc,selected,maths)
        notes={1:'Os ângulos são calculados por SPA no meio dos intervalos. As componentes locais GHI DNI e DHI são reconstruídas na exploração inversa. A variabilidade usa os pares disponíveis de cada dia.',
               2:'A reflexão do solo pertence aos cenários de transposição e não deve ser somada novamente à POA medida. A atenuação de sujidade é hipotética. A projeção de sombra é um exemplo geométrico sem dimensões das fileiras reais.',
               3:'A temperatura operacional aproxima a temperatura da célula. Curvas CEC e a associação 20 em série por 6 em paralelo não constituem validação de campo. A curva CC PVWatts é auxiliar. A POA reconstruída tem fechamento numérico sem validação independente das componentes.',
               4:'As perdas por estágio são calculadas e não identificam causas reais isoladamente. O PR exige a mesma janela de energia e irradiação. O fator empírico usa apenas abril. O balanço semitransparente BIPV é hipotético e seu SHGC equivalente não constitui certificação.'}
        doc.add_paragraph(notes[number])
    doc.add_heading('Convenções para a redação',level=1)
    doc.add_paragraph('Usar o mesmo plano, intervalo e máscara temporal nas comparações. G designa irradiância; H irradiação; P potência e E energia. '
                      'Cada linha define todas as variáveis da relação. i indexa intervalos ou observações; j indexa perdas ou estágios conforme a linha. '
                      'As operações soma, produto, média, valor absoluto, mínimo e máximo não são novas variáveis.')
    conventions=data['notation_conventions']
    selected_conventions={1:[1,2],2:[1,4],3:[2,3,4,5,6,7],4:[2,3,5,6,8,9]}
    indices=range(1,len(conventions)) if section is None else selected_conventions[section]
    for index in indices:doc.add_paragraph(conventions[index])
    doc.add_paragraph('Cada linha conta uma relação. Constantes e parâmetros não entram na contagem. N identifica relações do notebook e C, relações complementares do código; '
                      'selecionar a inclusão conforme a discussão e o espaço. A numeração final do capítulo será definida na consolidação.')
    doc.add_paragraph('O caso comum é PVDAQ 2107, abril a junho de 2024. As 4401 observações diurnas pareadas não representam todos os intervalos do trimestre. '
                      'A potência nominal declarada é 893 kWp. O exemplo de associação elétrica e o balanço BIPV são hipóteses identificadas.')
    doc.add_heading('Origem identificável das equações',level=1)
    groups={}
    for row in rows:groups.setdefault(row['source_locator'],[]).append(row['id'])
    for locator,ids in groups.items():doc.add_paragraph(locator+'  '+', '.join(ids))
    doc.add_paragraph('As fórmulas N correspondem às células de método das 19 subseções do notebook. '
                      'Os algoritmos de SPA, CEC e Perez–Driesse, as conversões e os filtros estão no código. Citar as fontes R e explicitar as hipóteses na redação.')
    doc.add_paragraph('Fontes e referências metodológicas R: '+REPOSITORY+'/blob/main/references/Referencias_do_Projeto.md')
    doc.add_paragraph('Inventário legível com localização e links: '+REPOSITORY+'/blob/main/docs/equacoes/README.md')
    path.parent.mkdir(parents=True,exist_ok=True);doc.save(path)


def access_links(section=None):
    """Offer document viewers before downloads; never open a DOCX blob as text."""
    name='Inventario_Geral_Equacoes_SPGER0046.docx' if section is None else f'Quadro_Equacoes_Secao_{section:02d}.docx'
    publication=json.loads((ROOT/'docs/auditoria/Publicacao_Drive_Equacoes.json').read_text())
    record=next(r for r in publication['files'] if Path(r['path']).name==name)
    links=[]
    if record.get('pdf_drive_url'):
        links.append('[Ler PDF no Drive]('+record['pdf_drive_url']+')')
    links.append('[Abrir Word no Drive]('+record['drive_url']+')')
    links.append('[Baixar Word](https://raw.githubusercontent.com/samuelasroma-byte/SPGER0046_Energias_Renovaveis/main/'+record['path']+')')
    if record.get('pdf_path'):
        links.append('[PDF no GitHub]('+REPOSITORY+'/blob/main/'+record['pdf_path']+')')
    return ' · '.join(links)


def write_catalogue(data):
    folder=ROOT/'docs/equacoes';folder.mkdir(parents=True,exist_ok=True)
    (folder/'Inventario_Equacoes.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
    fields=['id','section','subsection','origin','application_class','application','title','latex','function','symbols','source_locator','code_correspondence','refs','notes']
    with (folder/'Inventario_Equacoes.csv').open('w',encoding='utf-8-sig',newline='') as handle:
        writer=csv.DictWriter(handle,fields,extrasaction='ignore',lineterminator='\n');writer.writeheader()
        for row in data['equations']:writer.writerow(dict(row,refs='; '.join(row['refs'])))
    lines=['# Inventário de equações por seção e subseção','',
           'O levantamento contém **61 relações numeradas**: **37 escritas nas células de método do notebook** e **24 complementares explicitadas no código**. '
           'Cada ID tem uma única subseção de inserção. A equação de Faiman aparece uma vez em 3.2; em 2.1 há remissão a E3.03.','',
           'Uma linha representa uma relação numerada. Constantes e parâmetros não são novas equações. Os itens N formam o núcleo de inserção; '
           'os itens C documentam o método e sua inclusão depende do espaço e da discussão. Os documentos estão organizados por seção e subseção.','',
           '## Abrir os documentos','',
           'Para consultar as equações já diagramadas no celular, escolha **Ler PDF no Drive**. Para editar as equações nativas, escolha **Abrir Word no Drive**. '
           '**Baixar Word** entrega o arquivo para abrir no Word ou em outro editor compatível.','',
           '**Inventário geral**: '+access_links(),'',
           *['- **Seção '+str(number)+'**: '+access_links(number) for number in range(1,5)],'',
           '[Dados do inventário em CSV](Inventario_Equacoes.csv) · [Fonte estruturada em JSON](Inventario_Equacoes.json)','',
           '[Pasta com os cinco Word editáveis e os cinco PDF de leitura](https://drive.google.com/drive/folders/1yfwwArqBTVwPt7EYer9DlHL1vmeO5V7B).','',
           '| Seção | Notebook N | Código C | Total | Documento exclusivo |','| --- | --- | --- | --- | --- |']
    for number in range(1,5):
        rows=[r for r in data['equations'] if r['section']==number];cc=counts(rows)
        relative=f'../../secoes/{FOLDERS[number-1]}/equacoes/README.md'
        lines.append(f'| {number} {TITLES[number-1]} | {cc["N"]} | {cc["C"]} | {len(rows)} | [PDF e Word da seção {number}]({relative}) |')
    lines+=['| Total | 37 | 24 | 61 | |','',
            '## Correspondência com o cálculo do PVDAQ','',data['application_scope'],'',
            '**N/C identifica o registro da fórmula; a aplicação identifica sua função no estudo.** Uma expressão nas células de método pode ser fundamentação ou cenário. '
            'SPA, CEC e Perez–Driesse são calculados pelas bibliotecas; o inventário apresenta suas relações de redação, sem expandir todos os algoritmos internos.','',
            '| Aplicação verificada | Relações |','| --- | --- |']
    for label,count in Counter(r['application'] for r in data['equations']).items():lines.append(f'| {label} | {count} |')
    lines+=['','## Símbolos e unidades','',
            'Cada uma das 61 relações tem lista estruturada de símbolos, descrição e unidade no JSON e a mesma lista legível no Word e neste catálogo. '
            'Os índices também são identificados. O campo `code_correspondence` registra a operação, chamada ou coluna correspondente no código.','']
    for convention in data['notation_conventions']:lines+=['- '+convention]
    lines+=['',
            '## Fontes e convenções','',
            'As relações correspondem às células de método e às funções indicadas no código. As fórmulas são editáveis nos Word. '
            'As potências devem usar unidades coerentes; para integrar W em kWh, multiplicar pela duração em horas e dividir por 1000. '
            'A análise usa intervalos de 15 minutos e o período abril a junho de 2024 do PVDAQ 2107.','',
            'N corresponde às relações das células de método. C expõe relações utilizadas pelo código, inclusive métricas normalizadas, correlação, curva do inversor, ajuste empírico e balanço BIPV. '
            'As expressões internas de bibliotecas e os critérios computacionais não são expandidos como equações de redação.','']
    for sub in sorted({r['subsection'] for r in data['equations']}):
        rows=[r for r in data['equations'] if r['subsection']==sub]
        lines+=['## Subseção '+sub,'','| ID | Equação e função | Registro | Aplicação | Local verificável |','| --- | --- | --- | --- | --- |']
        for r in rows:
            path,locator=r['source_locator'].split('::')
            lines.append('| '+r['id']+' | **'+r['title']+'**. '+r['function']+' | '+r['origin']+' | '+r['application']+' | ['+locator+'](../../'+path+') |')
        lines.append('')
        for r in rows:lines+=['**'+r['id']+'**','', '$$'+r['latex']+'$$','',r['symbols'],'',
                            '**Correspondência no cálculo:** '+r['code_correspondence'],'',r['notes'],'']
    lines+=['## Referências metodológicas','',
            '[Fontes R do projeto](../../references/Referencias_do_Projeto.md). O código e o notebook permitem localizar cada relação; as fontes metodológicas devem ser citadas na redação final.','']
    (folder/'README.md').write_text('\n'.join(line.rstrip() for line in lines).rstrip()+'\n')
    return folder


def write_section_navigation(data):
    for number,name in enumerate(FOLDERS,1):
        folder=ROOT/'secoes'/name
        rows=[r for r in data['equations'] if r['section']==number]
        word=f'Quadro_Equacoes_Secao_{number:02d}.docx'
        text=f'# Equações da seção {number}\n\n'+access_links(number)+'\n\n[Inventário geral](../../../docs/equacoes/README.md)\n\n'
        text+='O PDF permite leitura; o Word contém as equações editáveis.\n\n'
        text+='Cada ID tem um único local de inserção. O Word inclui descrição e unidade de todos os símbolos de cada equação. N identifica fórmula escrita no notebook e C relação complementar do código; a aplicação distingue o cálculo principal das análises e cenários.\n\n'
        text+='| ID | Inserir em | Função | Registro | Aplicação |\n| --- | --- | --- | --- | --- |\n'
        for row in rows:text+=f'| {row["id"]} | {row["subsection"]} | {row["title"]} | {row["origin"]} | {row["application"]} |\n'
        if number==2:text+='\nA equação de Faiman é inserida em 3.2 como E3.03. Em 2.1 fazer remissão a esse ID.\n'
        (folder/'equacoes/README.md').write_text(text)
        path=folder/'README.md';text=path.read_text().split('## Equações da seção')[0].rstrip()
        path.write_text(text+'\n\n## Equações da seção\n\n[Ler as equações em PDF ou abrir o Word editável](equacoes/README.md).\n')
        path=folder/'documentos/README.md';text=path.read_text()
        old_link='[Quadro de equações exclusivo desta seção](../equacoes/'+word+').'
        text=text.replace(old_link,'').rstrip()
        link='[Quadro de equações exclusivo desta seção em PDF e Word](../equacoes/README.md).'
        if link not in text:path.write_text(text.rstrip()+'\n\n'+link+'\n')
        for path in folder.glob('subsecoes/*/README.md'):
            sub=path.parent.name.replace('_','.')
            selected=[r for r in rows if r['subsection']==sub];cc=counts(selected)
            text=path.read_text().split('## Inventário de equações')[0].rstrip()
            text+='\n\n## Inventário de equações\n\n[Quadro de equações da seção em PDF e Word](../../equacoes/README.md): '+str(cc['N'])+' relações do núcleo e '+str(cc['C'])+' de apoio nesta subseção. IDs: '+', '.join(r['id'] for r in selected)+'.\n'
            if sub=='2.1':text+='\nPara o modelo térmico de Faiman, citar E3.03 em 3.2 sem repetir sua inserção.\n'
            path.write_text(text)


def main():
    data=inventory();folder=write_catalogue(data)
    with tempfile.TemporaryDirectory(prefix='inventario_math_') as temporary:
        maths=math_elements(data['equations'],Path(temporary))
        paths=[folder/'Inventario_Geral_Equacoes_SPGER0046.docx']
        make_doc(data,maths,paths[0])
        for number in range(1,5):
            path=ROOT/'secoes'/FOLDERS[number-1]/f'equacoes/Quadro_Equacoes_Secao_{number:02d}.docx'
            make_doc(data,maths,path,number);paths.append(path)
    write_section_navigation(data)
    print(json.dumps({'relations':len(data['equations']),'origin':dict(counts(data['equations'])),
                      'section_counts':[sum(r['section']==i for r in data['equations']) for i in range(1,5)],
                      'DOCX':[str(p.relative_to(ROOT)) for p in paths]},ensure_ascii=False))


if __name__=='__main__':main()
