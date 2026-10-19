"""Extract editable tables/frames into four documents with exclusive owners.

Original consolidated DOCX files are kept intact. Numerical table XML is
copied; only the figure lists in the four planning frames are updated to the
current editorial allocation.
"""
from copy import deepcopy
from pathlib import Path
import json
import re
from docx import Document
from docx.table import Table
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def text(element):
    return ''.join(element.xpath('.//w:t/text()'))


def blocks(path, label):
    source = Document(path)
    result = {}
    current = None
    for element in source._element.body:
        if element.tag == qn('w:sectPr'):
            continue
        match = re.match(label + r' (\d+)\s*[—–-]', text(element)) if element.tag == qn('w:p') else None
        if match:
            current = int(match[1])
            result.setdefault(current, [])
        if current is None:
            continue
        # Blank manual page-break paragraphs can produce empty pages after extraction.
        if element.tag == qn('w:p') and not text(element).strip():
            continue
        result[current].append(deepcopy(element))
    return result


def build(root=ROOT):
    root = Path(root)
    policy = json.loads((root/'config/distribuicao_editorial.json').read_text())
    numerical = blocks(root/'outputs/tabelas/docx/Tabelas_do_Projeto.docx', 'Tabela')
    frames = blocks(root/'outputs/quadros/docx/Quadros_do_Projeto.docx', 'Quadro')
    figure_owners = {s: [f['id'].split('_')[0] for f in policy['figures'] if f['owner_subsection']==s]
                     for s in {f['owner_subsection'] for f in policy['figures']}}
    files = []
    for section in policy['sections']:
        number = str(section['section'])
        doc = Document(root/'outputs/tabelas/docx/Tabelas_do_Projeto.docx')
        for element in list(doc._element.body):
            if element.tag != qn('w:sectPr'):
                doc._element.body.remove(element)
        for style in doc.styles:
            for border in style.element.xpath('.//w:pBdr'):
                border.getparent().remove(border)
        doc.styles['Heading 1'].font.color.rgb = RGBColor(0,0,0)
        title = doc.add_paragraph('Tabelas e quadros da seção ' + number, 'Title')
        for run in title.runs:
            run.font.name = 'Times New Roman'; run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0,0,0); run.bold = True
        doc.add_paragraph(section['title'])
        doc.add_paragraph('Responsáveis pela seção: ' + section['responsible'] + '.')
        doc.add_paragraph('Inserir cada ID somente na subseção indicada; em outras partes, citar sem reproduzir. Numeração original preservada.')
        if number == '4':
            doc.add_paragraph('A Tabela 16 é um único item com quatro partes de continuação.')
        items = [dict(r,kind='T') for r in policy['word_tables']] + [dict(r,kind='Q') for r in policy['frames']]
        items = [r for r in items if r['owner_subsection'][0] == number]
        items.sort(key=lambda r:(tuple(map(int,r['owner_subsection'].split('.'))),r['kind']!='T',r['number']))
        previous = None
        for item in items:
            if item['owner_subsection'] != previous:
                p = doc.add_heading('Subseção ' + item['owner_subsection'], level=1)
                if previous is not None:
                    p.paragraph_format.page_break_before = True
                previous = item['owner_subsection']
            marker = doc.add_paragraph(item['id'] + '  Inserção exclusiva na subseção ' + item['owner_subsection'])
            marker.paragraph_format.keep_with_next = True
            for run in marker.runs:
                run.bold = True; run.font.size = Pt(12); run.font.color.rgb=RGBColor(0,0,0)
            elements = numerical[item['number']] if item['kind']=='T' else frames[item['number']]
            seen_caption = False
            inserted = []
            for original in elements:
                element = deepcopy(original)
                if item['id']=='T16' and element.tag==qn('w:p') and re.match(r'Tabela 16 ',text(element)):
                    if seen_caption:
                        from docx.oxml import OxmlElement
                        prop = element.find(qn('w:pPr'))
                        if prop is None:
                            prop=OxmlElement('w:pPr');element.insert(0,prop)
                        prop.append(OxmlElement('w:pageBreakBefore'))
                    seen_caption = True
                doc._element.body.insert(len(doc._element.body)-1,element)
                inserted.append(element)
                if item['kind']=='Q' and item['number'] >= 3 and element.tag == qn('w:tbl'):
                    table = Table(element, doc._body)
                    for row in table.rows[1:]:
                        match = re.match(r'([1-4]\.\d+)', row.cells[0].text)
                        if match:
                            cell = row.cells[1]
                            replacement = 'Figuras: ' + ', '.join(figure_owners.get(match[1],[])) + '.'
                            cell.text = re.sub(r'Figuras:[^\n]*', replacement, cell.text)
                            for p in cell.paragraphs:
                                p.paragraph_format.first_line_indent=Cm(0)
                                p.paragraph_format.space_after=Pt(0)
                                p.paragraph_format.line_spacing=1
                                for run in p.runs:
                                    run.font.name='Times New Roman';run.font.size=Pt(12)
                                    run.font.color.rgb=RGBColor(0,0,0)
            # Keep each short table/frame, its caption and notes on the same page.
            # T16 has explicit continuation blocks instead of one oversized table.
            from docx.text.paragraph import Paragraph
            groups = [[]]
            for element in inserted:
                if item['id']=='T16' and element.tag==qn('w:p') and re.match(r'Tabela 16 ',text(element)) and groups[-1]:
                    groups.append([])
                groups[-1].append(element)
            for group in groups:
                paragraphs=[]
                for element in group:
                    if element.tag==qn('w:p'):
                        paragraphs.append(Paragraph(element,doc._body))
                    else:
                        paragraphs.extend(Paragraph(p,doc._body) for p in element.xpath('.//w:p'))
                for p in paragraphs[:-1]:
                    p.paragraph_format.keep_with_next=True
                if paragraphs:
                    paragraphs[-1].paragraph_format.keep_with_next=False
        doc.core_properties.title='Tabelas e quadros da seção '+number
        doc.core_properties.author='Equipe SPGER0046 PPGER UFPB'
        path=root/section['directory']/f'tabelas/docx/Materiais_Tabulares_Secao_{int(number):02d}.docx'
        path.parent.mkdir(parents=True,exist_ok=True);doc.save(path);files.append(path)
        print(path.relative_to(root),len(doc.tables),'blocos de tabela',flush=True)
    return files


if __name__=='__main__':
    build()
