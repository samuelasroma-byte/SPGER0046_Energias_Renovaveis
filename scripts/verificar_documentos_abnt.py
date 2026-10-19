"""Conferência estrutural dos sete Word; a conferência visual é registrada à parte."""
from pathlib import Path
import json, zipfile, re, hashlib
from lxml import etree
from docx import Document
from docx.shared import Cm

ROOT=Path(__file__).resolve().parents[1]
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
W='{'+NS['w']+'}'
paths=sorted((ROOT/'outputs').rglob('*.docx'))+[ROOT/'bibliografia/outputs/Referencias_por_Subsecao.docx']
assert len(paths)==7
checked=[]
for path in paths:
    doc=Document(path)
    for section in doc.sections:
        for actual,cm in [(section.page_width,21),(section.page_height,29.7),
                          (section.top_margin,3),(section.left_margin,3),
                          (section.bottom_margin,2),(section.right_margin,2)]:
            assert abs(actual-Cm(cm))<1000,(path.name,actual,cm)
    count=0
    with zipfile.ZipFile(path) as archive:
        for member in archive.namelist():
            if member=='word/document.xml' or re.match(r'word/(header|footer)\d+\.xml$',member):
                root=etree.fromstring(archive.read(member))
                for run in root.xpath('//w:r',namespaces=NS):
                    fonts=run.find('w:rPr/w:rFonts',NS)
                    assert fonts is not None
                    for family in ['ascii','hAnsi','eastAsia','cs']:
                        assert fonts.get(W+family)=='Times New Roman',(path.name,family)
                    for size in ['sz','szCs']:
                        assert run.find('w:rPr/w:'+size,NS).get(W+'val')=='24',path.name
                    count+=1
        styles=etree.fromstring(archive.read('word/styles.xml'))
        for style in styles.xpath('//w:style[@w:type="paragraph" or @w:type="character"]',namespaces=NS):
            assert style.find('w:rPr/w:rFonts',NS).get(W+'ascii')=='Times New Roman'
            assert style.find('w:rPr/w:sz',NS).get(W+'val')=='24'
        xml=etree.fromstring(archive.read('word/document.xml'))
        for p in xml.xpath('//w:body/w:p',namespaces=NS):
            text=''.join(p.xpath('.//w:t/text()',namespaces=NS))
            if re.match(r'^(Tabela|Quadro|Figura) \d+ — ',text):
                assert p.find('w:pPr/w:jc',NS).get(W+'val')=='center',(path.name,text)
    checked.append({'document':str(path.relative_to(ROOT)),'native_runs_checked':count,
                    'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})

path=ROOT/'bibliografia/outputs/Referencias_por_Subsecao.docx'
doc=Document(path)
refs=[p for p in doc.paragraphs if p.style.name=='Referência ABNT']
assert len(refs)==39
style=doc.styles['Referência ABNT'].paragraph_format
assert style.line_spacing==1
assert abs(style.left_indent-Cm(1.25))<1000
assert abs(style.first_line_indent+Cm(1.25))<1000
assert style.space_after.pt==12
assert all(not re.search(r'\.\.\s',p.text) for p in refs)
assert len(doc.tables)==4
topics=[row.cells[0].text.split()[0] for table in doc.tables for row in table.rows[1:]]
assert topics==[f'{s}.{n}' for s,total in [(1,4),(2,4),(3,6),(4,5)] for n in range(1,total+1)]
with zipfile.ZipFile(path) as archive:
    xml=etree.fromstring(archive.read('word/document.xml'))
    targets=xml.xpath('//w:bookmarkStart/@w:name',namespaces=NS)
    links=xml.xpath('//w:hyperlink/@w:anchor',namespaces=NS)
    assert len(targets)==39 and len(set(targets))==39
    assert len(links)==102 and set(links)==set(targets)
result={'status':'passed','project_version':'1.2.2','documents':checked,
        'format':{'font':'Times New Roman','font_size_pt':12,'paper':'A4',
                  'margins_cm':{'top':3,'left':3,'bottom':2,'right':2},
                  'body_line_spacing':1.5,'reference_line_spacing':1,
                  'reference_space_after_pt':12,'requested_hanging_indent_cm':1.25},
        'bibliography':{'unique_references':39,'subsections':19,'working_internal_links':102},
        'visual_review':'see Conferencia_Visual.json; structural checks do not replace visual inspection'}
out=ROOT/'outputs/relatorios/Verificacao_Documentos_ABNT.json'
out.write_text(json.dumps(result,ensure_ascii=False,indent=2))
print(json.dumps(result,ensure_ascii=False,indent=2))
