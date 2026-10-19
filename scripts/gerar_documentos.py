"""Five native DOCX deliverables, with numerical tables and text frames separated."""
from pathlib import Path
import json
import math
import sys
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'bibliografia'))
import formatacao_abnt as abnt
save_document=abnt.save_document
OUT=ROOT/'outputs'
CSV=OUT/'tabelas/csv'
CFG=json.loads((ROOT/'config/projeto.json').read_text())
FIGURES=json.loads((OUT/'relatorios/catalogo_figuras.json').read_text())
SECTIONS=json.loads((ROOT/'docs/secoes_capitulo.json').read_text())
HAS_REVERSE=bool(json.loads((OUT/'resumo_execucao.json').read_text()).get('reverse_POA',{}).get('converged_hours',0))
FONT='Times New Roman'
SOURCE='Fonte: elaboração da equipe a partir de dados do PVDAQ, sistema 2107 (2026).'


def value(x,digits=2):
    if x is None or (isinstance(x,float) and math.isnan(x)):return 'Não definido'
    if isinstance(x,(int,float)):return f'{x:.{digits}f}'.replace('.',',')
    return str(x)


def paragraph(doc,text,bold=False,keep=False,center=False):
    p=doc.add_paragraph();p.paragraph_format.keep_with_next=keep
    compact=bool(bold or center or text.startswith(('Fonte:','Nota:','Referências metodológicas:')))
    if compact:
        p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1
        p.paragraph_format.space_before=Pt(6 if bold else 0);p.paragraph_format.space_after=Pt(6)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    r=p.add_run(text);r.bold=bold;r.font.name=FONT;r.font.size=Pt(12);r.font.color.rgb=RGBColor(0,0,0)
    return p


def new_document(title,opening):
    doc=Document()
    for element in list(doc.styles._element.iter(qn('w:pBdr'))):
        element.getparent().remove(element)
    for element in list(doc._element.iter(qn('w:pBdr'))):
        element.getparent().remove(element)
    abnt.configure_document(doc)
    for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
        st=doc.styles[name];st.font.name=FONT;st.font.color.rgb=RGBColor(0,0,0)
        st.font.size=Pt(12)
    doc.styles['Title'].font.size=Pt(12);doc.styles['Title'].font.bold=True
    p=doc.add_paragraph(title,'Title');p.paragraph_format.keep_with_next=True
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    paragraph(doc,opening)
    doc.core_properties.title=title;doc.core_properties.author='Grupo SPGER0046 UFPB'
    doc.core_properties.subject='Farm Solar Array PVDAQ 2107 abril a junho de 2024'
    return doc


def border(element,name,val='single',size=12):
    b=OxmlElement('w:'+name);b.set(qn('w:val'),val);b.set(qn('w:sz'),str(size));b.set(qn('w:color'),'000000')
    element.append(b)


def table(doc,title,headers,rows,widths=None,grid=False,note='',source=SOURCE):
    if len(rows)>30:
        chunks=[rows[i:i+23] for i in range(0,len(rows),23)]
        for i,chunk in enumerate(chunks):
            if i:doc.add_page_break()
            last=i==len(chunks)-1
            label='conclusão' if last else 'continua' if i==0 else 'continuação'
            result=table(doc,title+' ('+label+')',headers,chunk,widths,grid,
                         note if last else '',source if last else '')
        return result
    paragraph(doc,title,True,True,True)
    t=doc.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    props=t._tbl.tblPr
    old=props.find(qn('w:tblBorders'))
    if old is not None:props.remove(old)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        enabled=grid or edge in ['top','bottom']
        border(borders,edge,'single' if enabled else 'nil',6 if grid else 12)
    props.append(borders)
    if widths is None:widths=[16/len(headers)]*len(headers)
    factor=16/sum(widths);widths=[w*factor for w in widths]
    for column,width in zip(t.columns,widths):column.width=Cm(width)
    header=t.rows[0]
    rep=OxmlElement('w:tblHeader');header._tr.get_or_add_trPr().append(rep)
    for cell,text in zip(header.cells,headers):cell.text=str(text)
    for row in rows:
        cells=t.add_row().cells
        for cell,text in zip(cells,row):cell.text=str(text)
    for row_index,row in enumerate(t.rows):
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for col,cell in enumerate(row.cells):
            cell.width=Cm(widths[col]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr()
            cb=cp.find(qn('w:tcBorders'))
            if cb is not None:cp.remove(cb)
            cb=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                enabled=grid or (row_index==0 and edge in ['top','bottom']) or (row_index==len(t.rows)-1 and edge=='bottom')
                border(cb,edge,'single' if enabled else 'nil',6 if grid else 12)
            cp.append(cb)
            margins=OxmlElement('w:tcMar')
            for edge,n in [('top',65),('bottom',65),('left',80),('right',80)]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(n));e.set(qn('w:type'),'dxa');margins.append(e)
            cp.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                p.paragraph_format.line_spacing=1
                p.paragraph_format.first_line_indent=Cm(0)
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if (grid or col==0) else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.name=FONT;r.font.size=Pt(12);r.bold=row_index==0;r.font.color.rgb=RGBColor(0,0,0)
    if source:
        paragraph(doc,source,keep=bool(note))
        for cell in t.rows[-1].cells:
            for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
    if note:
        paragraph(doc,'Nota: '+note).paragraph_format.keep_together=True
    return t


def load(name):return pd.read_csv(CSV/(name+'.csv'))


def numerical_tables():
    doc=new_document('Tabelas do projeto fotovoltaico',
        'Farm Solar Array PVDAQ 2107, abril a junho de 2024. Resultados numéricos para a redação do capítulo. '
        'As tabelas são editáveis, com três traços horizontais e sem grade no corpo. Todos os indicadores informam sua janela de cálculo.')
    s=load('estatistica_descritiva')
    rows=[[f'{r.variable} ({r.unit})',str(int(r.valid)),value(r.coverage_percent),value(r['mean']),value(r['median']),value(r.std_sample)] for _,r in s.iterrows()]
    table(doc,'Tabela 1 — Estatística descritiva de posição e dispersão',
        ['Variável e unidade','n','Cobertura (%)','Média','Mediana','DP'],rows,[5.1,1.6,2.8,2.4,2.4,2.4],
        note='Mesmos 4.401 pares diurnos; 4.421 intervalos esperados pelo critério solar. DP amostral com ddof=1. A cobertura é conjunta, não a cobertura individual de cada sensor.')
    doc.add_page_break()
    table(doc,'Tabela 2 — Extremos e quartis dos mesmos pares',
        ['Variável e unidade','Mínimo','Q1','Q3','Máximo','IQR'],
        [[f'{r.variable} ({r.unit})',*[value(r[k]) for k in ['minimum','Q1','Q3','maximum','IQR']]] for _,r in s.iterrows()],
        [5.0,2.0,2.0,2.0,2.0,2.0],note='IQR=Q3−Q1. A moda não foi incluída para variáveis contínuas.')
    doc.add_page_break()
    m=load('metrics')
    labels=['Físico versus medidor','Físico versus inversores','Ajuste empírico em abril','Empírico verificado em maio','Empírico testado em junho']
    table(doc,'Tabela 3 — Erros de potência nas comparações principais',
        ['Comparação','n','MBE (kW)','MAE (kW)','RMSE (kW)'],
        [[labels[i],str(int(r.n)),value(r.MBE,3),value(r.MAE,3),value(r.RMSE,3)] for i,(_,r) in enumerate(m.iterrows())],
        [6,1.5,2.6,2.6,2.6],note='Erro=simulado−observado. O fator empírico foi ajustado somente em abril; o modelo físico permanece separado. R27.')
    table(doc,'Tabela 4 — Erros normalizados correlação e determinação',
        ['Comparação','nMBE (%)','nRMSE (%)','r','R²'],
        [[labels[i],value(r.nMBE_percent,3),value(r.nRMSE_percent,3),value(r.r,5),value(r.R2,5)] for i,(_,r) in enumerate(m.iterrows())],
        [6,2.5,2.5,2.2,2.2],note='Normalização pela média observada dos pares da respectiva comparação. R²=1−SSE/SST; não foi substituído por r². R27.')
    doc.add_page_break()
    q=load('metricas_por_mes');months={4:'Abril',5:'Maio',6:'Junho'}
    table(doc,'Tabela 5 — Erros mensais do modelo físico',
        ['Mês','n','Cobertura (%)','MBE (kW)','MAE (kW)','RMSE (kW)'],
        [[months[int(r.month)],str(int(r.n)),value(r.coverage_percent),value(r.MBE,3),value(r.MAE,3),value(r.RMSE,3)] for _,r in q.iterrows()],
        [2,1.6,2.5,2.7,2.7,2.7],note='Somente o modelo físico original; mesmos pares dentro de cada mês.')
    table(doc,'Tabela 6 — Erros relativos mensais do modelo físico',
        ['Mês','nMBE (%)','nRMSE (%)','r','R²'],
        [[months[int(r.month)],value(r.nMBE_percent,3),value(r.nRMSE_percent,3),value(r.r,5),value(r.R2,5)] for _,r in q.iterrows()],
        [3,3,3,3,3],note='Denominador: média da potência observada de cada mês. Médias e amostras diferem entre meses.')
    con=load('consistencia_medidor_inversores').iloc[0]
    table(doc,'Tabela 7 — Consistência entre soma dos inversores e medidor',
        ['n','MBE (kW)','MAE (kW)','RMSE (kW)','nRMSE (%)'],
        [[str(int(con.n)),value(con.MBE,3),value(con.MAE,3),value(con.RMSE,3),value(con.nRMSE_percent,3)]],
        [2,3,3,3,3],note='Erro definido como soma dos inversores−medidor. Esta tabela compara medições entre si; não é uma simulação.')
    doc.add_page_break()
    q=load('energia_e_indicadores_mensais')
    table(doc,'Tabela 8 — Energia e cobertura dos intervalos pareados',
        ['Mês','n','Cobertura (%)','Medida (MWh)','Simulada (MWh)'],
        [[months[int(r.month)],str(int(r.n_paired)),value(r.coverage_percent),value(r.E_observed_paired_kWh/1000,3),value(r.E_modeled_paired_kWh/1000,3)] for _,r in q.iterrows()],
        [2.3,1.5,3,3.5,3.5],note='Integração de médias de potência ×0,25 h. Janela diurna pareada.')
    table(doc,'Tabela 9 — Energia observada integral por mês',
        ['Mês','Registros','Esperados','Energia medida (MWh)'],
        [[months[int(r.month)],str(int(r.meter_full_month_records)),str(int(r.meter_full_month_expected)),value(r.E_observed_full_month_kWh/1000,3)] for _,r in q.iterrows()],
        [3,3.2,3.2,5],note='Total do medidor para o mês inteiro, separado do recorte usado nos gráficos de validação. Não utilizar este total com POA parcial para PR.')
    doc.add_page_break()
    y=load('yields').iloc[0]
    table(doc,'Tabela 10 — Produtividades e desempenho na janela pareada',
        ['Indicador','Valor','Unidade'],
        [['Yr',value(y.Yr_h,3),'h'],['Ya modelada',value(y.Ya_model_h,3),'h'],
         ['Yf simulada',value(y.Yf_model_kWh_kWp,3),'kWh/kWp'],['Yf observada',value(y.Yf_observed_kWh_kWp,3),'kWh/kWp'],
         ['PR simulado',value(100*y.PR_model,3),'%'],['PR observado pareado',value(100*y.PR_observed_paired,3),'%'],
         ['Lc modelada',value(y.Lc_model_h,3),'h'],['Ls modelada',value(y.Ls_model_h,3),'h']],
        [8,3.5,3.5],note='P₀=893 kWp declarados. Yr, Ya, Yf e PR usam a mesma máscara de 4.401 intervalos. Ya não é medição DC independente. R12 e R30.')
    stage=load('energy_stages')
    table(doc,'Tabela 11 — Energia por estágio do modelo',
        ['Estágio','Energia (MWh)'],
        [[r.stage,value(r.energy_kWh/1000,3)] for _,r in stage.iterrows()],[11,4],
        note='Decomposição calculada; os incrementos ou perdas entre estágios não identificam sozinhos causas reais. R08, R11 e R12.')
    doc.add_page_break()
    sen=load('environment_sensitivity')
    table(doc,'Tabela 12 — Sensibilidade de sujidade e ventilação',
        ['Sujidade assumida (%)','Condição térmica','Energia AC (MWh)'],
        [[value(100*r.soiling_assumed,0),r.thermal_scenario,value(r.E_AC_paired_kWh/1000,3)] for _,r in sen.iterrows()],
        [4,7.5,4],note='Atenuação adicional do módulo relativamente à leitura POA. Não é perda de sujeira medida. Mesmos pares da cadeia principal. R07 e R28.')
    cap=load('capacity_sensitivity')
    table(doc,'Tabela 13 — Sensibilidade do limite AC por inversor',
        ['Limite assumido','MBE (kW)','RMSE (kW)','nRMSE (%)'],
        [[r.comparison.replace('Paco = ',''),value(r.MBE,3),value(r.RMSE,3),value(r.nRMSE_percent,3)] for _,r in cap.iterrows()],
        [6.5,3,3,3],note='Comparação de hipóteses; não selecionadas pelo melhor resultado de junho. R13.')
    if HAS_REVERSE and (CSV/'environment_scenarios.csv').exists():
        doc.add_page_break();q=load('environment_scenarios')
        table(doc,'Tabela 14 — Albedo e geometrias com componentes estimadas',
            ['Cenário','H POA estimada (kWh/m²)'],
            [[r.scenario,value(r.H_POA_kWh_m2,3)] for _,r in q.iterrows()],[10,5],
            note='Mesmas 1.037 horas elegíveis da inversão; não é irradiação integral do trimestre. Sensibilidade de albedo mantendo componentes estimadas com referência 0,20. R14–R16 e R29.')
        q=load('bipv_comparison')
        table(doc,'Tabela 15 — Produtividade e PR dos cenários BIPV',
            ['Geometria','Yf (kWh/kWp)','PR (%)','Horas'],
            [[r.scenario,value(r.E_AC_for_1kWp_kWh,3),value(100*r.PR_modeled,3),str(int(r.n_input_hours))] for _,r in q.iterrows()],
            [6,4,3,2],note='Cenários verticais com menor ventilação. Mesma base temporal e capacidade de normalização, sem medições de fachada. R14 e R19.')
    doc.add_page_break()
    q=load('energia_e_HSP_diurna_pareada')
    table(doc,'Tabela 16 — Energia diária e horas equivalentes na janela pareada',
        ['Dia de 2024','Medida (MWh)','Simulada (MWh)','PSH (h)','n','Cobertura (%)'],
        [[pd.Timestamp(r.day_local).strftime('%d/%m'),value(r.E_measured_kWh/1000,3),
          value(r.E_modeled_kWh/1000,3),value(r.PSH_POA_paired_h,3),str(int(r.paired_intervals)),value(r.coverage_percent,1)] for _,r in q.iterrows()],
        [2.2,3.1,3.1,2.5,1.4,3.0],note='PSH=H_POA/(1 kW/m²). Somente intervalos diurnos pareados, não duração de brilho solar nem total integral do dia. n esperado e datas com fuso estão no CSV. R12 e R30.')
    dest=OUT/'tabelas/docx/Tabelas_do_Projeto.docx';save_document(doc,dest)
    return dest


def textual_frames():
    doc=new_document('Quadros do projeto fotovoltaico',
        'Definições, dados, hipóteses e ligação às 19 subseções do capítulo. Os quadros são editáveis, com grade completa. '
        'A natureza da evidência diferencia medição, cálculo e cenário; fundamentação bibliográfica e redação final permanecem a cargo do grupo.')
    data=pd.read_csv(OUT/'quadros/situacao_dos_dados.csv')
    table(doc,'Quadro 1 — Dados disponíveis e limites do estudo',list(data.columns),data.values.tolist(),[5.5,4.3,7],True,
        'Dados PVDAQ 2107. Metadados incompletos não foram preenchidos como se fossem medições.')
    doc.add_page_break()
    assumptions=[
        ['Geometria','25°; azimute 180°','Metadados; faltam dimensões locais das fileiras.'],
        ['Capacidade declarada','893 kWp','Referência para Yf e PR; placa soma 892,8 kWp e CEC 891,648 kWp.'],
        ['Temperatura','Faiman U₀=25; U₁=6,84','Coeficientes genéricos; aproximação de célula sem medição própria.'],
        ['Sujidade','0%, 2% e 5%','Cenários de atenuação adicional à leitura POA, sem inferência de perda real.'],
        ['Perdas DC','Fiação 2%; mismatch 2%; conexões 0,5%','Hipóteses multiplicativas; não obtidas por medição individual.'],
        ['Albedo','0,20; sensibilidade 0,10 e 0,40','Hipóteses para a transposição exploratória, sem ganho bifacial atribuído.'],
        ['Inversores','24; 30 kW de limite assumido por unidade','Sensibilidade alternativa a 27,6 kW; configuração instalada não confirmada.'],
        ['Strings','Exemplo de 20 série × 6 paralelo','Equivalente de 120 módulos por grupo, não diagrama real verificado.'],
        ['BIPV semitransparente','100 m²; ρ=0,12; τsolar=0,25; τvisível=0,35','Propriedades hipotéticas; SHGC equivalente simplificado, sem certificação.'],
    ]
    table(doc,'Quadro 2 — Parâmetros documentados e hipóteses do modelo',
        ['Parâmetro','Valor e natureza','Implicação'],assumptions,[3.8,5.3,7.7],True,
        'As unidades dos coeficientes térmicos são W/(m² K) e W s/(m³ K), respectivamente. R07, R10, R13 e R20.')
    titles={1:'Fundamentos e recurso solar',2:'Ambiente e instalação',3:'Modelagem da conversão',4:'Desempenho e validação'}
    for number in range(1,5):
        doc.add_page_break()
        rows=[]
        for s in SECTIONS:
            if s['id'].startswith(str(number)+'.'):
                rows.append([s['id']+' '+s['title']+'\n'+s['owner'],
                    s['outputs']+'\nEvidência: '+s['evidence']+'.\nFiguras: '+', '.join(x.split('_')[0] for x in s['figures'])+'.',
                    s['limitation']+'\nFontes: '+', '.join(s['refs'])+'.'])
        table(doc,f'Quadro {number+2} — {titles[number]}',
            ['Subseção e responsável','Entregas e evidência','Limite e referências'],rows,[5.9,5.9,6.0],True,
            'A matriz CSV e Markdown contém também objetivos, entradas e métodos. Cobertura didática não significa validação independente de cada etapa.')
    dest=OUT/'quadros/docx/Quadros_do_Projeto.docx';save_document(doc,dest);return dest


def figure_document(category,filename,title,opening,include_all=False):
    items=FIGURES if include_all else [r for r in FIGURES if r['category']==category]
    doc=new_document(title,opening)
    for i,r in enumerate(items,1):
        caption=paragraph(doc,f'Figura {i} — {r["title"]}',True,True,True)
        if i>1:caption.paragraph_format.page_break_before=True
        path=ROOT/r['files']['claro']['png']
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(5)
        p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        width=15 if category=='figuras' and i==1 else 16
        pic=p.add_run().add_picture(str(path),width=Cm(width))
        pic._inline.docPr.set('descr',r['caption'])
        caption_text=paragraph(doc,r['caption'])
        caption_text.paragraph_format.first_line_indent=Cm(0);caption_text.paragraph_format.line_spacing=1
        caption_text.paragraph_format.space_after=Pt(6);caption_text.alignment=WD_ALIGN_PARAGRAPH.LEFT
        paragraph(doc,SOURCE)
        paragraph(doc,'Referências metodológicas: '+', '.join(r['refs'])+'.')
    dest=OUT/category/'docx'/filename;save_document(doc,dest);return dest


def main():
    files=[numerical_tables(),textual_frames(),
        figure_document('graficos','Graficos_do_Projeto.docx','Gráficos e diagramas do projeto fotovoltaico',
            'Coleção completa das 30 figuras em fundo branco para a redação: 25 gráficos, um fluxograma e quatro diagramas. '
            'Dados e modelos são identificados nas legendas. Os originais estão em PDF vetorial e PNG a 600 dpi nos temas claro e escuro.',True),
        figure_document('fluxogramas','Fluxogramas_do_Projeto.docx','Fluxogramas do projeto fotovoltaico',
            'Fluxo da cadeia principal e das extensões exploratórias. As figuras incorporadas têm títulos e legendas editáveis; os originais PDF e PNG estão nas pastas correspondentes.'),
        figure_document('figuras','Ilustracoes_FV.docx','Ilustrações do sistema fotovoltaico',
            'Esquemas científicos do plano do arranjo, componentes radiativas, organização elétrica e extensão BIPV. Os exemplos hipotéticos estão identificados nas legendas.')]
    for p in files:print(str(p),p.stat().st_size,flush=True)
    return files


if __name__=='__main__':main()
