"""Guia de uso por tópico e catálogo alfabético de referências completas."""
from collections import defaultdict
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import formatacao_abnt as abnt

SOURCE='Fonte: elaboração da equipe a partir de metadados bibliográficos do Crossref e do OpenAlex (2026).'


def compact(doc,text,bold=False,keep=False,center=False):
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0)
    p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(12)
    p.paragraph_format.keep_with_next=keep
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.add_run(text).bold=bold
    return p


def frame(doc,caption,topics,survey):
    compact(doc,caption,True,True,True)
    table=doc.add_table(rows=1,cols=3)
    table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    widths=[8.0,3.2,4.8]
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'6');x.set(qn('w:color'),'000000');borders.append(x)
    table._tbl.tblPr.append(borders)
    for col,width in zip(table.columns,widths):col.width=Cm(width)
    header=table.rows[0]
    for cell,title in zip(header.cells,['Subseção e assunto','Responsável','Obras indicadas']):cell.text=title
    repeat=OxmlElement('w:tblHeader');header._tr.get_or_add_trPr().append(repeat)
    for topic in topics:
        cells=table.add_row().cells
        cells[0].text=topic['id']+' '+topic['title'];cells[1].text=topic['owner']
        p=cells[2].paragraphs[0]
        works=[w for w in survey.selected if w['topic']==topic['id']]
        for i,w in enumerate(works):
            if i:p.add_run(', ')
            abnt.internal_link(p,w['reference_code'],w['reference_code'])
    for i,row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell,width in zip(row.cells,widths):
            cell.width=Cm(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margin=OxmlElement('w:tcMar')
            for edge,value in [('top',80),('bottom',80),('left',80),('right',80)]:
                x=OxmlElement('w:'+edge);x.set(qn('w:w'),str(value));x.set(qn('w:type'),'dxa');margin.append(x)
            cell._tc.get_or_add_tcPr().append(margin)
            for p in cell.paragraphs:
                f=p.paragraph_format;f.first_line_indent=Cm(0);f.line_spacing=1
                f.space_before=Pt(0);f.space_after=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                if i==0:
                    for run in p.runs:run.bold=True
    compact(doc,SOURCE)
    return table


def write(survey):
    doc=abnt.configure_document(Document())
    doc.add_paragraph('REFERÊNCIAS PARA A FUNDAMENTAÇÃO DO CAPÍTULO FOTOVOLTAICO','Title')
    doc.add_paragraph('Este guia relaciona as obras indicadas para as quatro seções e as 19 subseções do capítulo SPGER0046 da UFPB. Os quadros mostram o assunto, o responsável e os códigos das obras. A lista alfabética apresenta cada referência completa uma única vez, seguida das subseções em que ela pode ser utilizada.')
    doc.add_paragraph('Os códigos B01 a B39 são links internos para localizar as referências. Os códigos R, quando apresentados na lista, correspondem às fontes metodológicas identificadas nos resultados do projeto. A equipe deve conferir o texto integral e as páginas pertinentes antes de redigir as citações.')
    doc.add_paragraph('O levantamento selecionou 39 obras a partir de 829 registros deduplicados. As buscas executadas utilizaram Crossref e OpenAlex. Contagens de citações, critérios e registros das consultas permanecem nos catálogos CSV, separados deste guia de leitura.')
    names={1:'Fundamentos e recurso solar',2:'Ambiente e instalação',3:'Modelagem da conversão',4:'Desempenho e validação'}
    for n in range(1,5):
        if n>1:doc.add_page_break()
        frame(doc,f'Quadro {n} — Referências da seção {n} {names[n]}',
              [t for t in survey.topics if t['id'].startswith(str(n)+'.')],survey)
    doc.add_page_break()
    p=doc.add_paragraph('REFERÊNCIAS','Heading 1');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    compact(doc,'Lista alfabética das obras indicadas nos quadros. Os códigos identificam as obras para consulta; a referência bibliográfica é o parágrafo que começa pela autoria.')
    use=defaultdict(set)
    for w in survey.selected:use[w['id']].add(w['topic'])
    unique={w['id']:w for w in survey.selected}
    for number,w in enumerate(sorted(unique.values(),key=abnt.sort_key),1):
        code=w['reference_code']
        kind='Livro' if w['document_type']=='livro' else 'Relatório técnico' if 'report' in str(w['document_type']) else 'Trabalho em evento' if 'proceeding' in str(w['document_type']) else 'Artigo' if 'article' in str(w['document_type']) else 'Documento técnico'
        p=compact(doc,'REFERÊNCIA '+code+' — '+kind,True,True)
        p.paragraph_format.space_before=Pt(12);p.paragraph_format.space_after=Pt(6)
        abnt.bookmark(p,code,number)
        p=abnt.add_reference(doc,w);p.paragraph_format.keep_with_next=True
        p=compact(doc,'',keep=False);p.paragraph_format.keep_together=True
        p.add_run('Aplicação no capítulo: ').bold=True
        p.add_run('subseções '+', '.join(sorted(use[w['id']],key=lambda v:tuple(map(int,v.split('.')))))+'.')
        legacy=[s for s in w['seed_ids'] if s.startswith('R')]
        if legacy:p.add_run(' Código nos resultados: '+', '.join(legacy)+'.')
    doc.core_properties.author='Equipe SPGER0046 UFPB'
    doc.core_properties.title='Referências para a fundamentação do capítulo fotovoltaico'
    doc.core_properties.subject='Referências por subseção com apresentação ABNT e recuo francês solicitado'
    return abnt.save_document(doc,survey.out/'Referencias_por_Subsecao.docx')
