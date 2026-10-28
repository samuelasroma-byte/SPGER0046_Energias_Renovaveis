"""Padronização editorial dos manuscritos sem recalcular os resultados científicos.

Corpo 12, fontes 10, equações OMML com SEQ à direita, figuras SVG/PNG 600 dpi
e quadros semânticos centralizados. O PPTX coletivo não é modificado.
"""
from pathlib import Path
from copy import deepcopy
import json,re,subprocess,hashlib,io,zipfile
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.part import Part
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import fitz
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/'_trabalho_editorial'/'padronizacao'
FIG=ROOT/'figuras'
ASVG='http://schemas.microsoft.com/office/drawing/2016/SVG/main'

DEFINICOES={
 'zenital':r'em que $\theta_z$ é o ângulo zenital geométrico, medido a partir da vertical local; $\phi$ é a latitude geográfica, positiva ao norte; $\delta$ é a declinação solar, relativa ao equador celeste; e $\omega$ é o ângulo horário, nulo no meio-dia solar, negativo pela manhã e positivo à tarde. Os ângulos são apresentados em graus; nas funções trigonométricas do código, são convertidos para radianos.',
 'azimute':r'em que $A_s$ é o azimute solar contado do norte no sentido horário, no intervalo $0^{\circ}\leq A_s<360^{\circ}$; $\gamma_s$ é o azimute contado do sul, negativo a leste e positivo a oeste, conforme a convenção do livro; e o operador $\mathrm{mod}\;360^{\circ}$ reduz o resultado a esse intervalo. Ambos os ângulos são expressos em graus. A relação não define um azimute no zênite.',
 'orbital':r'em que $G_{0n}(t)$ é a irradiância extraterrestre normal no instante $t$, em W/m²; $G_{\mathrm{sc}}$ é a irradiância de referência na distância $r_0$, também em W/m²; $r_0$ é a distância de referência Terra–Sol; e $r(t)$ é a distância Terra–Sol no instante considerado. As duas distâncias devem ter a mesma unidade, de modo que sua razão e seu quadrado sejam adimensionais.',
 'spencer':r'em que $E_0$ é o fator orbital adimensional e $B=2\pi(n-1)/365$ é o argumento angular em radianos, sendo $n$ o número do dia no ano. Os coeficientes numéricos são adimensionais. A irradiância normal é calculada por $G_{0n}=G_{\mathrm{sc}}E_0$, com $G_{\mathrm{sc}}=1366{,}1\,\mathrm{W/m^2}$. O denominador 365 é mantido em 2024 para reproduzir a implementação conferida do pvlib.',
 'projecao':r'em que $G_{0h}$ é a irradiância extraterrestre projetada na horizontal e $G_{0n}$ é a irradiância extraterrestre normal, ambas em W/m²; $\theta_z$ é o ângulo zenital geométrico; e $\max(0,\cos\theta_z)$ seleciona o maior entre zero e o cosseno do zenital. Esse fator é adimensional e anula a projeção quando o centro do Sol está abaixo do horizonte geométrico.',
 'integral':r'em que $H_{[t_1,t_2]}$ é a irradiação acumulada no plano considerado entre os limites $t_1$ e $t_2$; $G(t)$ é a irradiância nesse plano, em W/m²; $t$ é a variável temporal; e $\mathrm{d}t$ é o elemento de tempo da integração. Com o tempo em segundos, $H$ é obtida em J/m²; com o tempo em horas, em Wh/m².',
 'soma':r'em que $H$ é a irradiação dos intervalos incluídos; $\overline{G}_i$ é a irradiância média do intervalo $i$, em W/m²; $\Delta t_i$ é sua duração; e $\mathcal{V}$ é o conjunto de intervalos válidos sobre os quais se realiza a soma. O índice $i$ identifica cada intervalo incluído. Se a duração é dada em horas, o resultado é Wh/m²; se é dada em segundos, é J/m². A soma não inclui a energia de intervalos ausentes.',
 'fechamento':r'em que GHI é a irradiância global horizontal, DNI é a irradiância direta normal e DHI é a irradiância difusa horizontal, todas em W/m²; $\theta_z$ é o ângulo zenital; e $\mathrm{DNI}\cos\theta_z$ é a contribuição direta projetada na horizontal. A relação pressupõe o Sol acima do horizonte e definições compatíveis de banda espectral, campo angular, plano e suporte temporal.'
}
ORDEM=['zenital','azimute','orbital','spencer','projecao','integral','soma','fechamento']

def digest(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()

def paragrafo_md(text):
    WORK.mkdir(parents=True,exist_ok=True)
    key=hashlib.sha256(text.encode()).hexdigest()[:16];md=WORK/(key+'.md');out=WORK/(key+'.docx')
    if not out.exists():
        md.write_text(text)
        subprocess.run(['pandoc',str(md),'-f','markdown+tex_math_dollars','-t','docx','-o',str(out)],check=True,capture_output=True)
    return deepcopy(Document(out).paragraphs[0]._p)

def inserir_depois(anchor,text):
    xml=paragrafo_md(text);anchor._p.addnext(xml)
    return Paragraph(xml,anchor._parent)

def formatar_texto(p,size=12,align=None):
    p.paragraph_format.first_line_indent=Cm(1.25)
    p.paragraph_format.line_spacing=1.5;p.paragraph_format.space_after=Pt(6)
    p.paragraph_format.widow_control=True
    if align is not None:p.alignment=align
    for r in p.runs:
        r.font.name='Times New Roman';r.font.size=Pt(size);r.font.color.rgb=RGBColor(0,0,0);r.italic=False
    for mr in p._p.xpath('.//m:r'):
        pr=mr.find(qn('w:rPr'))
        if pr is None:pr=OxmlElement('w:rPr');mr.insert(0,pr)
        mr.remove(pr);mr.insert(1 if len(mr) and mr[0].tag==qn('m:rPr') else 0,pr)
        for k in ['sz','szCs']:
            node=pr.find(qn('w:'+k))
            if node is None:node=OxmlElement('w:'+k);pr.append(node)
            node.set(qn('w:val'),str(round(size*2)))

def fontes_caption(p,title=False):
    formatar_texto(p,10,WD_ALIGN_PARAGRAPH.CENTER if title else WD_ALIGN_PARAGRAPH.LEFT)
    pf=p.paragraph_format;pf.first_line_indent=Cm(0);pf.left_indent=Cm(.4);pf.right_indent=Cm(.4)
    pf.line_spacing=1;pf.space_before=Pt(8 if title else 3);pf.space_after=Pt(4 if title else 8)
    pf.keep_with_next=title
    for r in p.runs:r.bold=title

def assets():
    """Deriva a versão clara do diagrama já conferido e preserva a G01 vetorial."""
    FIG.mkdir(exist_ok=True)
    source=FIG/'D01_geometria.tex'
    if not source.exists():
        from preparar_slides import assets as preparar
        preparar()
    tex=source.read_text().replace(r'\definecolor{fg}{HTML}{FFFFFF}',r'\definecolor{fg}{HTML}{000000}').replace(r'\definecolor{bg}{HTML}{000000}',r'\definecolor{bg}{HTML}{FFFFFF}').replace('{FFD45C}','{8B5700}').replace('{70DCF5}','{00728C}')
    stem=FIG/'Manuscrito_Figura01_Geometria';stem.with_suffix('.tex').write_text(tex)
    r=subprocess.run(['xelatex','-interaction=nonstopmode','-halt-on-error','-output-directory',str(FIG),str(stem.with_suffix('.tex'))],capture_output=True,text=True)
    if r.returncode:raise RuntimeError(r.stdout[-2000:])
    target=FIG/'Manuscrito_Figura02_PVDAQ'
    target.with_suffix('.pdf').write_bytes((FIG/'G01_recurso_geometria_claro.pdf').read_bytes())
    # Os dois SVG são derivados dos PDF vetoriais. Essa conversão também evita
    # diferenças de recorte das legendas do SVG original entre leitores Word.
    for i,st in enumerate([stem,target],1):
        subprocess.run(['pdftocairo','-svg',str(st.with_suffix('.pdf')),str(st.with_suffix('.svg'))],check=True,capture_output=True)
        svg=st.with_suffix('.svg').read_text()
        identifiers=re.findall(r'\bid="([^"]+)"',svg)
        for identifier in sorted(set(identifiers),key=len,reverse=True):
            svg=svg.replace('id="'+identifier+'"','id="fig'+str(i)+'_'+identifier+'"').replace('#'+identifier+'"','#fig'+str(i)+'_'+identifier+'"').replace('#'+identifier+')','#fig'+str(i)+'_'+identifier+')')
        st.with_suffix('.svg').write_text(svg)
    for st in [stem,target]:
        with fitz.open(st.with_suffix('.pdf')) as d:
            pix=d[0].get_pixmap(matrix=fitz.Matrix(3600/d[0].rect.width,3600/d[0].rect.width),alpha=False)
            image=Image.open(io.BytesIO(pix.tobytes('png'))).convert('RGB')
            image.save(st.with_suffix('.png'),dpi=(600,600))
    return [stem,target]

def substituir_por_vetor(doc,p,stem,title):
    p.clear();p.paragraph_format.first_line_indent=Cm(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(0);p.paragraph_format.keep_with_next=True
    shape=p.add_run().add_picture(str(stem.with_suffix('.png')),width=Cm(15.2))
    shape._inline.docPr.set('descr',title);shape._inline.docPr.set('title',title)
    partname=doc.part.package.next_partname('/word/media/figura_vetorial%d.svg')
    svgpart=Part(partname,'image/svg+xml',stem.with_suffix('.svg').read_bytes(),doc.part.package)
    rid=doc.part.relate_to(svgpart,RT.IMAGE)
    blip=shape._inline.xpath('.//a:blip')[0]
    extlst=OxmlElement('a:extLst');ext=OxmlElement('a:ext');ext.set('uri','{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
    node=OxmlElement('w:dummy');node.tag='{'+ASVG+'}svgBlip';node.set(qn('r:embed'),rid)
    ext.append(node);extlst.append(ext);blip.append(extlst)

def figures(doc,stems):
    pics=[p for p in doc.paragraphs if p._p.xpath('.//w:drawing')]
    if not pics:return []
    assert len(pics)==2
    titles=['Figura 1 – Geometria solar local e convenção azimutal','Figura 2 – POA observada e ângulos solares no PVDAQ 2107']
    sources=['Fonte: elaborado pelo autor, com base em Duffie, Beckman e Blair (2020).','Fonte: dados PVDAQ 2107 (Deline et al., 2021); elaboração do autor (2026).']
    result=[]
    for i,p in enumerate(pics):
        after=p._p.getnext()
        if after is not None and after.tag==qn('w:p'):
            nxt=Paragraph(after,p._parent)
            if nxt.text.startswith(('Corte vertical','Figura G01','Figura 1','Figura 2')):after.getparent().remove(after)
        title=p.insert_paragraph_before(titles[i]);fontes_caption(title,True)
        substituir_por_vetor(doc,p,stems[i],titles[i])
        src=inserir_depois(p,sources[i]);fontes_caption(src)
        if i==0:
            body=title.insert_paragraph_before('A Figura 1 reúne o corte vertical e a vista em planta dessas convenções.')
            formatar_texto(body,12,WD_ALIGN_PARAGRAPH.JUSTIFY);body.paragraph_format.keep_with_next=True
        im=Image.open(stems[i].with_suffix('.png'))
        result.append({'figura':i+1,'titulo':titles[i],'vetor':stems[i].with_suffix('.svg').name,'png_px':list(im.size),'largura_inserida_cm':15.2,'dpi_efetivo_horizontal':im.width/(15.2/2.54),'svg_sha256':digest(stems[i].with_suffix('.svg'))})
    # Os detalhes metodológicos permanecem no corpo, não na linha de fonte.
    for p in doc.paragraphs:
        for node in p._p.xpath('.//w:t'):
            if node.text:
                node.text=node.text.replace('Figura G01','Figura 2').replace('conforme a coleção gráfica identificada como G01','conforme a Figura 2')
        if p.text.startswith('O gráfico reúne naturezas distintas de informação.'):
            p.add_run(' Na Figura 2, o zenital é geométrico e a incidência utiliza o zenital aparente, conforme a cadeia computacional vigente.')
    return result

def field(p,instruction,result):
    node=OxmlElement('w:fldSimple');node.set(qn('w:instr'),instruction)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');size=OxmlElement('w:sz');size.set(qn('w:val'),'24');pr.append(size);r.append(pr)
    t=OxmlElement('w:t');t.text=str(result);r.append(t);node.append(r);p._p.append(node)

def equation(p,n,key):
    math=deepcopy(p._p.xpath('.//m:oMath')[0])
    for child in list(math):
        if re.search(r'E1\.0[1-5]',''.join(t.text or '' for t in child.iter(qn('m:t')))):math.remove(child)
    p.clear();p.style='Normal';p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    pf=p.paragraph_format;pf.first_line_indent=Cm(0);pf.left_indent=Cm(0);pf.right_indent=Cm(0)
    pf.line_spacing=1.15;pf.space_before=Pt(8);pf.space_after=Pt(8);pf.keep_together=True;pf.keep_with_next=True
    pf.tab_stops.clear_all();pf.tab_stops.add_tab_stop(Cm(8),WD_TAB_ALIGNMENT.CENTER);pf.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT)
    p.add_run('\t');p._p.append(math);p.add_run('\t(')
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(700+n));start.set(qn('w:name'),'Equacao_'+key);p._p.append(start)
    field(p,' SEQ Equacao \\* ARABIC ',n)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(700+n));p._p.append(end);p.add_run(')')
    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(12)
    for mr in math.iter(qn('m:r')):
        pr=mr.find(qn('w:rPr'))
        if pr is None:pr=OxmlElement('w:rPr');mr.insert(0,pr)
        mr.remove(pr);mr.insert(1 if len(mr) and mr[0].tag==qn('m:rPr') else 0,pr)
        size=pr.find(qn('w:sz'))
        if size is None:size=OxmlElement('w:sz');pr.append(size)
        size.set(qn('w:val'),'24')

def numbering(doc,slug):
    displays=[p for p in doc.paragraphs if p._p.xpath('.//m:oMathPara')]
    keys=ORDEM if slug=='Secao_01_Integrada' else ORDEM[:5] if slug.startswith('1_1') else ORDEM[5:] if slug.startswith('1_2') else []
    assert len(displays)==len(keys),(slug,len(displays),keys)
    nums={key:n for n,key in enumerate(keys,1)}
    legacy={'E1.01':nums.get('zenital'),'E1.02':nums.get('projecao'),'E1.03':nums.get('integral'),'E1.04':nums.get('soma'),'E1.05':nums.get('fechamento')}
    for p in doc.paragraphs:
        for t in p._p.xpath('.//w:t'):
            if t.text:
                for old,n in legacy.items():
                    if n:t.text=t.text.replace('('+old+')','('+str(n)+')').replace(old,str(n))
    result=[]
    for n,(p,key) in enumerate(zip(displays,keys),1):
        before=Paragraph(p._p.getprevious(),p._parent)
        old_after=p._p.getnext()
        # As definições são deslocadas para depois da equação, sem repeti-las.
        short_leads={
          'zenital':f'A relação geométrica esférica entre a posição do observador e a direção solar é apresentada na equação ({n}):',
          'azimute':f'A conversão entre a convenção azimutal de Duffie, Beckman e Blair (2020, p. 13) e a convenção do SPA é dada pela equação ({n}):',
          'spencer':f'O fator orbital da aproximação de Spencer implementada no pvlib 0.16.1 é expresso pela equação ({n}) (PVLIB CONTRIBUTORS, 2026):',
          'integral':f'A irradiação recebida no plano e intervalo considerados é obtida pela integração temporal da irradiância, conforme a equação ({n}):'
        }
        if key in short_leads:before.text=short_leads[key]
        # Conserva o contexto já escrito e altera apenas sua chamada editorial.
        for t in before._p.xpath('.//w:t'):
            if not t.text:continue
            s=t.text
            if key=='zenital':s=s.replace('A relação geométrica esférica elementar é ('+str(n)+'):',f'A relação geométrica esférica elementar é apresentada na equação ({n}):')
            elif key=='azimute':s=s.replace(' é:',f' é dada pela equação ({n}):')
            elif key=='orbital':s=s.replace('Uma representação conceitual dessa dependência é:',f'Essa dependência é representada pela equação ({n}):')
            elif key=='spencer':s=s.replace('O fator empregado é (',f'O fator empregado é expresso pela equação ({n}) (')
            elif key=='projecao':s=s.replace('('+str(n)+'):',f'na equação ({n}):')
            elif key=='integral':s=re.sub(r'\bé\s*$',f'é dada pela equação ({n}):',s)
            elif key=='soma':s=re.sub(r'\bé\s*$',f'é expressa pela equação ({n}):',s)
            elif key=='fechamento':s=re.sub(r'\bé\s*$',f'é expresso pela equação ({n}):',s)
            t.text=s
        assert 'equação ('+str(n)+')' in before.text,(key,before.text)
        before.paragraph_format.keep_with_next=True
        markers={'zenital':'A elevação solar','orbital':'A relação explicita','projecao':'Durante o dia','integral':'A equivalência','soma':'A forma discreta é aqui derivada','fechamento':'A equação deve utilizar'}
        if key in markers and old_after is not None and old_after.tag==qn('w:p'):
            marker=markers[key];oldp=Paragraph(old_after,p._parent)
            assert marker in oldp.text,(key,'interpretação seguinte não localizada')
            for child in list(old_after):
                if child.tag==qn('w:pPr'):continue
                nodes=list(child.iter(qn('w:t')))
                hit=next((t for t in nodes if marker in (t.text or '')),None)
                if hit is not None:
                    hit.text=hit.text[hit.text.index(marker):];break
                old_after.remove(child)
        equation(p,n,key)
        definition=inserir_depois(p,DEFINICOES[key]);definition.style='Body Text'
        formatar_texto(definition,12,WD_ALIGN_PARAGRAPH.JUSTIFY)
        result.append({'numero':n,'conceito':key,'definicao_dos_termos':True,'chamada_anterior':before.text,'posicao_do_numero':'tabulação direita a 16 cm','representacao':'OMML editável + campo SEQ Equacao'})
    return result

def quadros(doc):
    results=[]
    for i,tab in enumerate(doc.tables,1):
        first=tab.rows[0].cells[0].text
        title='Símbolos, grandezas e convenções adotadas' if first=='Símbolo' else 'Grandezas radiométricas e planos de referência'
        source='Fonte: elaborado pelo autor, com base em Duffie, Beckman e Blair (2020).' if first=='Símbolo' else 'Fonte: elaborado pelo autor, com base em Duffie, Beckman e Blair (2020) e Sengupta et al. (2024).'
        before=tab._tbl.getprevious()
        if before is not None and before.tag==qn('w:p') and Paragraph(before,tab._parent).text.lower().startswith('quadro de conceitos'):
            before.getparent().remove(before)
        caption=OxmlElement('w:p');tab._tbl.addprevious(caption);cp=Paragraph(caption,tab._parent);cp.add_run(f'Quadro {i} – {title}');fontes_caption(cp,True)
        prev=caption.getprevious()
        if prev is not None and prev.tag==qn('w:p'):
            pp=Paragraph(prev,tab._parent)
            if f'Quadro {i}' not in pp.text:pp.add_run(f' O Quadro {i} sintetiza essas grandezas e convenções.')
        after=tab._tbl.getnext()
        if after is not None and after.tag==qn('w:p') and Paragraph(after,tab._parent).text.startswith('Fonte:'):after.getparent().remove(after)
        source_xml=paragrafo_md(source);tab._tbl.addnext(source_xml);fontes_caption(Paragraph(source_xml,tab._parent))
        tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
        widths=[2.1,8.0,5.1] if first=='Símbolo' else [2.7,8.0,4.5]
        for col,w in zip(tab.columns,widths):col.width=Cm(w)
        props=tab._tbl.tblPr
        indent=props.find(qn('w:tblInd'))
        if indent is not None:props.remove(indent)
        width=props.find(qn('w:tblW'));width.set(qn('w:type'),'dxa');width.set(qn('w:w'),str(round(15.2/2.54*1440)))
        borders=props.find(qn('w:tblBorders'))
        if borders is None:borders=OxmlElement('w:tblBorders');props.append(borders)
        for edge in ['top','left','bottom','right','insideH','insideV']:
            node=borders.find(qn('w:'+edge))
            if node is None:node=OxmlElement('w:'+edge);borders.append(node)
            node.set(qn('w:val'),'single');node.set(qn('w:sz'),'8');node.set(qn('w:color'),'000000')
        for ri,row in enumerate(tab.rows):
            pr=row._tr.get_or_add_trPr()
            if pr.find(qn('w:cantSplit')) is None:pr.append(OxmlElement('w:cantSplit'))
            for ci,cell in enumerate(row.cells):
                cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=cell._tc.get_or_add_tcPr();margins=tcpr.find(qn('w:tcMar'))
                if margins is None:margins=OxmlElement('w:tcMar');tcpr.append(margins)
                for edge,val in [('top',70),('bottom',70),('left',100),('right',100)]:
                    node=OxmlElement('w:'+edge);node.set(qn('w:w'),str(val));node.set(qn('w:type'),'dxa');margins.append(node)
                for p in cell.paragraphs:
                    formatar_texto(p,12,WD_ALIGN_PARAGRAPH.CENTER)
                    p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(0)
                    p.paragraph_format.keep_with_next=True
                    for r in p.runs:r.bold=ri==0
            if ri==0 and pr.find(qn('w:tblHeader')) is None:pr.append(OxmlElement('w:tblHeader'))
        results.append({'quadro':i,'titulo':title,'largura_cm':15.2,'centralizado':True,'colunas_centralizadas':[1,2,3],'borda_pt':1,'fonte_pt':10})
    return results

def run(root=ROOT):
    root=Path(root);stems=assets();reports=[]
    for slug in ['1_1_Geometria_Solar','1_2_Grandezas_Radiometricas','1_3_Modulacao_Atmosferica','1_4_Medicao_e_Estimativa','Secao_01_Integrada']:
        path=root/'manuscritos'/(slug+'.docx');doc=Document(path)
        for name in ['Normal','Body Text','First Paragraph','Title','Heading 1','Heading 2','Heading 3']:
            if name in doc.styles:
                st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(12);st.font.italic=False
                if name in ['Normal','Body Text','First Paragraph']:st.paragraph_format.line_spacing=1.5
        for p in doc.paragraphs:
            formatar_texto(p,12)
            if p.style.name in ['Title','Heading 1','Heading 2','Heading 3']:
                p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True
            if p.text.startswith(('SPGER0046','Samuel Constantino Teixeira','Convenções:')):p.paragraph_format.first_line_indent=Cm(0)
            if p.text.startswith(('DELINE,','DRIESSE,','DUFFIE,','NOAA GLOBAL','PVLIB CONTRIBUTORS','REDA,','SENGUPTA,','TEIXEIRA,')):
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.line_spacing=1;p.paragraph_format.first_line_indent=Cm(-.6)
        f=figures(doc,stems);eq=numbering(doc,slug);q=quadros(doc)
        for section in doc.sections:
            for p in section.header.paragraphs:
                for r in p.runs:r.font.size=Pt(10)
        # Impede compressão posterior das imagens de compatibilidade no Word.
        settings=doc.settings.element
        if settings.find(qn('w:doNotAutoCompressPictures')) is None:settings.append(OxmlElement('w:doNotAutoCompressPictures'))
        live=set(doc.element.xpath('//@r:embed'))|set(doc.element.xpath('//@r:link'))
        for rid,rel in list(doc.part.rels.items()):
            if rel.reltype==RT.IMAGE and rid not in live:doc.part.drop_rel(rid)
        doc.save(path)
        reports.append({'arquivo':path.name,'corpo_pt':12,'fontes_ilustracoes_pt':10,'equacoes':eq,'figuras':f,'quadros':q,'sha256_docx':digest(path)})
    (root/'documentacao/Padronizacao_Editorial_Manuscritos.json').write_text(json.dumps(reports,ensure_ascii=False,indent=2))
    (root/'documentacao/Catalogo_Equacoes_Manuscrito.json').write_text(json.dumps({r['arquivo']:r['equacoes'] for r in reports},ensure_ascii=False,indent=2))
    return reports

if __name__=='__main__':
    print(json.dumps([{'arquivo':r['arquivo'],'equacoes':len(r['equacoes']),'figuras':len(r['figuras']),'quadros':len(r['quadros'])} for r in run()],ensure_ascii=False,indent=2))
